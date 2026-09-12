"""Ingest expert corrections and turn them into something that improves runs.

The loop: the tool writes a workbook with a **Verification** sheet listing each
extracted value with its verbatim callout, its grid reference and blank
*Correct value / Verified by / Date* columns. An engineer marks it up. This
module reads the marked-up file back and does three things with it.

1. **Score.** Per-field accuracy — how many values were right, how many wrong,
   what the wrong ones should have been. Tracked over time so a change can be
   shown to help rather than argued about.
2. **Teach by example.** A drawing whose extraction was confirmed correct
   becomes a verified few-shot example for sheets in the same family.
3. **Teach by grammar — the one that actually fixes misses.** When a corrected
   value came from a callout the regex grammar could not parse, that callout is
   reported as an unparsed pattern. Adding a pattern to `callout_grammar` fixes
   that shape of callout permanently, for every future drawing, with no model
   involvement at all.

What this is not: fine-tuning. That needs a 24 GB GPU and a hundred labelled
sheets, neither of which is on this machine, and pretending otherwise would
waste the corrections. The three mechanisms above are what a local setup can
genuinely learn from.
"""
from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from extractors import callout_grammar as G

VERIFY_SHEET = "Verification"
REPORT_DIR = Path("output/accuracy")

# Column positions in the Verification sheet, 1-based (see workbook_extras).
COL_ITEM, COL_TAG, COL_EXTRACTED = 2, 3, 4
COL_CALLOUT, COL_GRID, COL_CONF = 5, 6, 7
COL_CORRECT, COL_BY, COL_DATE = 8, 9, 10


@dataclass
class Correction:
    """One row an expert touched."""
    item: str
    tag: str
    extracted: str
    corrected: str
    callout: str = ""
    grid_ref: str = ""
    confidence: str = ""
    verified_by: str = ""
    verified_at: str = ""

    @property
    def changed(self) -> bool:
        return _norm(self.corrected) != "" and _norm(self.corrected) != _norm(self.extracted)


@dataclass
class SheetReview:
    """What one marked-up workbook says about one drawing."""
    drawing_no: str = ""
    source: str = ""
    rows: list[Correction] = field(default_factory=list)

    @property
    def reviewed(self) -> list[Correction]:
        """Rows an expert actually signed off — corrected or explicitly ticked."""
        return [r for r in self.rows
                if _norm(r.corrected) or r.verified_by.strip()]

    @property
    def wrong(self) -> list[Correction]:
        return [r for r in self.rows if r.changed]

    @property
    def right(self) -> list[Correction]:
        return [r for r in self.reviewed if not r.changed]

    def accuracy(self) -> float:
        n = len(self.reviewed)
        return (len(self.right) / n) if n else 0.0


def _norm(value) -> str:
    """Compare values the way a person would: 600 == 600.0 == '600 '."""
    if value is None:
        return ""
    text = str(value).strip()
    if not text:
        return ""
    try:
        return f"{float(text):g}"
    except ValueError:
        return re.sub(r"\s+", " ", text).lower()


def read_review(xlsx_path: str | Path) -> SheetReview:
    """Read a marked-up workbook's Verification sheet."""
    path = Path(xlsx_path)
    wb = load_workbook(path, data_only=True)
    if VERIFY_SHEET not in wb.sheetnames:
        raise ValueError(
            f"{path.name} has no '{VERIFY_SHEET}' sheet — corrections are read "
            f"from that sheet's 'Correct value' column. Re-generate the "
            f"workbook, or mark up the one the tool produced.")
    ws = wb[VERIFY_SHEET]

    review = SheetReview(source=str(path))
    header = str(ws.cell(row=2, column=1).value or "")
    m = re.search(r"Drawing:\s*(\S+)", header)
    if m:
        review.drawing_no = m.group(1)

    for row in ws.iter_rows(min_row=6, values_only=True):
        if not row or not isinstance(row[0], int):
            continue
        def get(idx):
            return row[idx - 1] if len(row) >= idx else None
        review.rows.append(Correction(
            item=str(get(COL_ITEM) or ""), tag=str(get(COL_TAG) or ""),
            extracted=str(get(COL_EXTRACTED) or ""),
            corrected=str(get(COL_CORRECT) or ""),
            callout=str(get(COL_CALLOUT) or ""),
            grid_ref=str(get(COL_GRID) or ""),
            confidence=str(get(COL_CONF) or ""),
            verified_by=str(get(COL_BY) or ""),
            verified_at=str(get(COL_DATE) or "")))
    return review


def unparsed_callouts(review: SheetReview) -> list[dict]:
    """Corrected rows whose callout the grammar cannot read.

    This is the highest-value output of the whole loop. A callout the grammar
    misses is missed on every drawing, for ever, until a pattern is added — and
    unlike model behaviour, a regex fix is deterministic and testable.
    """
    out = []
    for row in review.wrong:
        if not row.callout.strip():
            continue
        if row.item.strip().lower() != "pedestal":
            continue
        if G.parse_pedestal_callout(row.callout) is None:
            out.append({"drawing_no": review.drawing_no, "tag": row.tag,
                        "callout": row.callout, "expected": row.corrected,
                        "grid_ref": row.grid_ref})
    return out


def score(reviews: list[SheetReview]) -> dict:
    """Aggregate accuracy across marked-up workbooks."""
    per_item: dict[str, dict[str, int]] = {}
    per_drawing: dict[str, dict] = {}
    for review in reviews:
        stats = per_drawing.setdefault(
            review.drawing_no or Path(review.source).stem,
            {"reviewed": 0, "right": 0, "wrong": 0})
        stats["reviewed"] += len(review.reviewed)
        stats["right"] += len(review.right)
        stats["wrong"] += len(review.wrong)
        for row in review.reviewed:
            bucket = per_item.setdefault(row.item or "(unlabelled)",
                                         {"reviewed": 0, "right": 0, "wrong": 0})
            bucket["reviewed"] += 1
            bucket["wrong" if row.changed else "right"] += 1

    total = sum(s["reviewed"] for s in per_drawing.values())
    right = sum(s["right"] for s in per_drawing.values())
    return {
        "generated_at": date.today().isoformat(),
        "drawings_reviewed": len(per_drawing),
        "values_reviewed": total,
        "values_correct": right,
        "accuracy_pct": round(100 * right / total, 1) if total else 0.0,
        "per_item": {k: {**v, "accuracy_pct": round(100 * v["right"] / v["reviewed"], 1)
                         if v["reviewed"] else 0.0}
                     for k, v in sorted(per_item.items())},
        "per_drawing": per_drawing,
    }


def ingest(paths: list[str | Path], *, examples_dir: Path | None = None,
           write_report: bool = True) -> dict:
    """Read marked-up workbooks, score them, and bank what can be learned."""
    from extractors import rag_examples

    reviews, failures = [], []
    for path in paths:
        try:
            reviews.append(read_review(path))
        except (ValueError, OSError) as exc:
            failures.append(f"{Path(path).name}: {exc}")

    report = score(reviews)
    report["failures"] = failures

    patterns: list[dict] = []
    saved_examples: list[str] = []
    for review in reviews:
        patterns.extend(unparsed_callouts(review))

        # A drawing whose reviewed pedestal rows were all confirmed correct is
        # worth keeping as a few-shot example. One wrong row disqualifies it:
        # an example is treated as ground truth by every later run, so a
        # half-right one actively teaches the mistake.
        peds = [r for r in review.reviewed if r.item.strip().lower() == "pedestal"]
        if peds and not any(r.changed for r in peds) and review.drawing_no:
            rows = []
            for r in peds:
                parsed = G.parse_pedestal_callout(r.callout)
                if parsed:
                    rows.append({"tag": parsed["tag"],
                                 "length_mm": parsed["length_mm"],
                                 "width_mm": parsed["width_mm"],
                                 "quantity": parsed["quantity"]})
            if rows:
                who = next((r.verified_by for r in peds if r.verified_by.strip()), "")
                path = rag_examples.save_example(
                    drawing_no=review.drawing_no, revision="", image_hash="",
                    verified_extraction={"pedestals": rows},
                    verified_by=who or "expert review",
                    source_pdf=Path(review.source).name,
                    notes="Confirmed correct during expert review.",
                    examples_dir=examples_dir)
                saved_examples.append(path.name)

    report["unparsed_callouts"] = patterns
    report["examples_saved"] = saved_examples

    if write_report:
        REPORT_DIR.mkdir(parents=True, exist_ok=True)
        out = REPORT_DIR / f"accuracy_{date.today().isoformat()}.json"
        out.write_text(json.dumps(report, indent=2), encoding="utf-8")
        report["report_path"] = str(out)
    return report
