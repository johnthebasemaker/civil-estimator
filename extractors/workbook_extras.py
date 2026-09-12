"""Extra sheets appended to the standard 10-sheet workbook.

Written against the finished file with openpyxl rather than inside
`core/excel_writer.py`, so the mandated 10-sheet template keeps its exact
identity and Phase 1 stays untouched.

  * **Verification** — one row per extracted value with its verbatim callout and
    drawing grid reference, and blank columns for sign-off. This is what a civil
    team marks up when checking the takeoff against the drawing.
  * **Extraction_Log** — every model call, timing and confidence note. The audit
    trail for how a number got into the BOQ.
  * **Derivation** — each derived quantity with the formula that produced it, so
    an assumption is arguable rather than buried.
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from extractors.models import ExtractionResult
from extractors.sheet_grid import DEFAULT_GRID, SheetGrid
from extractors.verification import VERIFY_HEADERS, verification_rows

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=14, color="C00000")
ENTRY_FILL = PatternFill("solid", fgColor="FFF2CC")     # columns to be filled in
LEFT = Alignment(horizontal="left", vertical="top", wrap_text=True)


def _header_row(ws, row: int, headers: list[str]) -> None:
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=col, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center",
                                wrap_text=True)


def _widths(ws, widths: dict[int, int]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w


def append_verification_sheet(xlsx_path: str | Path, result: ExtractionResult,
                              grid: SheetGrid = DEFAULT_GRID) -> Path:
    """The sheet a civil team marks up against the drawing."""
    path = Path(xlsx_path)
    wb = load_workbook(path)
    if "Verification" in wb.sheetnames:
        del wb["Verification"]
    ws = wb.create_sheet("Verification")

    ws["A1"] = "EXTRACTION VERIFICATION — check each row against the drawing"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = (f"Drawing: {result.title_block.drawing_no or '-'}  Rev "
                f"{result.title_block.revision or '-'}   ·   "
                f"read by {result.model}   ·   "
                f"{datetime.now().strftime('%Y-%m-%d %H:%M')}")
    ws["A3"] = ("Grid ref locates the value on the sheet border (e.g. D-7). "
                "Fill the shaded columns. Anything wrong: write the correct "
                "value and it goes back into the estimate.")
    ws["A3"].alignment = LEFT
    for r in (2, 3):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=11)

    head = 5
    _header_row(ws, head, VERIFY_HEADERS)
    rows = verification_rows(result, grid)
    for i, row in enumerate(rows, start=1):
        for col, val in enumerate(row, start=1):
            c = ws.cell(row=head + i, column=col, value=val)
            c.alignment = LEFT
            if col in (8, 9, 10):          # correct value / verified by / date
                c.fill = ENTRY_FILL

    total = head + len(rows) + 2
    ws.cell(row=total, column=1, value="Checked by").font = Font(bold=True)
    ws.cell(row=total, column=4, value="Signature").font = Font(bold=True)
    ws.cell(row=total, column=6, value="Date").font = Font(bold=True)
    for col in (2, 5, 7):
        ws.cell(row=total, column=col).fill = ENTRY_FILL

    _widths(ws, {1: 5, 2: 14, 3: 10, 4: 30, 5: 42, 6: 12, 7: 11,
                 8: 22, 9: 16, 10: 12, 11: 40})
    ws.freeze_panes = f"A{head + 1}"
    wb.save(path)
    return path


def append_derivation_sheet(xlsx_path: str | Path, derived: list) -> Path:
    """Every derived quantity with the arithmetic that produced it."""
    path = Path(xlsx_path)
    if not derived:
        return path
    wb = load_workbook(path)
    if "Derivation" in wb.sheetnames:
        del wb["Derivation"]
    ws = wb.create_sheet("Derivation")

    ws["A1"] = "DERIVED QUANTITIES — not read from the drawing"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = ("These follow from the slab geometry plus site convention. "
                "Check the assumptions below before pricing.")
    ws.merge_cells("A2:E2")

    head = 4
    _header_row(ws, head, ["#", "Rule", "Element", "Tag", "How it was worked out"])
    for i, item in enumerate(derived, start=1):
        ws.cell(row=head + i, column=1, value=i)
        ws.cell(row=head + i, column=2, value=item.rule_key)
        ws.cell(row=head + i, column=3, value=item.target_field)
        ws.cell(row=head + i, column=4, value=getattr(item.element, "tag", ""))
        c = ws.cell(row=head + i, column=5, value=item.explanation)
        c.alignment = LEFT
    _widths(ws, {1: 5, 2: 14, 3: 18, 4: 14, 5: 80})
    ws.freeze_panes = f"A{head + 1}"
    wb.save(path)
    return path


def append_audit_sheet(xlsx_path: str | Path, result: ExtractionResult,
                       plan=None) -> Path:
    """Every model call, timing and note — how the numbers got here."""
    path = Path(xlsx_path)
    wb = load_workbook(path)
    if "Extraction_Log" in wb.sheetnames:
        del wb["Extraction_Log"]
    ws = wb.create_sheet("Extraction_Log")
    bold = Font(bold=True)
    row = 1

    ws.cell(row=row, column=1, value="VISION EXTRACTION AUDIT LOG").font = TITLE_FONT
    row += 2

    def put(label: str, value: str) -> None:
        nonlocal row
        ws.cell(row=row, column=1, value=label).font = bold
        c = ws.cell(row=row, column=2, value=value)
        c.alignment = LEFT
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
        row += 1

    put("Source PDF", Path(result.source_pdf).name)
    put("Page", str(result.page_number + 1))
    put("Model", result.model)
    put("Strategy / profile", result.profile)
    put("Sheet dHash", result.image_hash)
    put("Text blocks located", f"{result.text_blocks_found} (no model calls)")
    put("Montages sent", str(result.montages_sent))
    put("Elapsed", f"{result.total_elapsed_s:.0f} s")
    put("Generated", datetime.now().strftime("%Y-%m-%d %H:%M"))
    put("Status", "UNVERIFIED DRAFT — requires engineer review")
    row += 1

    ws.cell(row=row, column=1, value="Extracted pedestals").font = bold
    row += 1
    _header_row(ws, row, ["Tag", "L (mm)", "W (mm)", "Nos", "Confidence",
                          "Grid ref", "Verbatim callout"])
    row += 1
    for p in result.pedestals:
        for col, val in enumerate([p.tag, p.length_mm, p.width_mm, p.quantity,
                                   p.confidence, p.grid_ref, p.raw_text], start=1):
            ws.cell(row=row, column=col, value=val)
        row += 1
    row += 1

    ws.cell(row=row, column=1, value="Model calls").font = bold
    row += 1
    _header_row(ws, row, ["Stage", "Region", "Pixels", "Seconds", "Status"])
    row += 1
    for r in result.responses:
        status = "skipped (blank)" if r.skipped else (r.error or "ok")
        for col, val in enumerate([r.stage, r.region,
                                   f"{r.px_width}x{r.px_height}",
                                   r.elapsed_s, status], start=1):
            ws.cell(row=row, column=col, value=val)
        row += 1
    row += 1

    ws.cell(row=row, column=1, value="Notes").font = bold
    row += 1
    notes = list(result.confidence_notes) + list(getattr(plan, "warnings", []) or [])
    for n in notes:
        c = ws.cell(row=row, column=1, value=f"• {n}")
        c.alignment = LEFT
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        row += 1

    _widths(ws, {1: 26, 2: 22, 3: 14, 4: 10, 5: 14, 6: 12, 7: 46})
    wb.save(path)
    return path


def append_discovery_sheet(xlsx_path: str | Path, result: ExtractionResult) -> Path:
    """Everything the drawing specifies, in the drawing's own words.

    The ten standard sheets are organised around element types chosen when the
    template was designed. A drawing is under no obligation to contain them: on
    a sections-and-details sheet there is no pedestal to put on the Pedestals
    tab and no slab for the Slabs tab, yet the sheet is dense with grout beds,
    coating, bar callouts and anchor assemblies that all cost money.

    So this sheet carries whatever `extractors.discovery` found, named by the
    drawing. A blank quantity is not a failure — it is the sheet stating a
    specification without an extent, and the shaded cell is where the reviewing
    engineer supplies the missing area, length or count.
    """
    path = Path(xlsx_path)
    disc = result.discovery or {}
    items = disc.get("items") or []
    specs = disc.get("specifications") or []
    heights = disc.get("heights") or []
    if not (items or specs or heights):
        return path

    wb = load_workbook(path)
    if "Drawing_Items" in wb.sheetnames:
        del wb["Drawing_Items"]
    ws = wb.create_sheet("Drawing_Items")

    ws["A1"] = "READ FROM THE DRAWING — named by the sheet, not by a preset list"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = (f"{result.title_block.drawing_no or Path(result.source_pdf).stem} · "
                f"{len(items)} item(s), {disc.get('measured', 0)} with a quantity "
                f"the drawing states outright. Shaded cells need a figure from "
                f"you before the row can be priced.")
    ws["A2"].alignment = LEFT
    ws.merge_cells("A2:H2")

    head = 4
    _header_row(ws, head, ["#", "Description", "UoM", "Qty stated", "Basis",
                           "Seen", "Grid ref", "Source text"])
    for i, item in enumerate(items, start=1):
        r = head + i
        values = [i, item.get("description"), item.get("uom"), item.get("qty"),
                  item.get("basis"), item.get("occurrences"),
                  item.get("grid_ref"), item.get("source")]
        for col, value in enumerate(values, start=1):
            c = ws.cell(row=r, column=col, value=value)
            if col in (2, 5, 8):
                c.alignment = LEFT
            if col == 4:
                c.number_format = "#,##0.###"
            if item.get("confirm") and col in (4, 5):
                c.fill = ENTRY_FILL

    r = head + len(items) + 2
    for row in specs:
        ws.cell(row=r, column=1, value="Specification")
        ws.cell(row=r, column=2, value=f"{row['value']:g} {row['unit']}")
        c = ws.cell(row=r, column=3, value=row["text"])
        c.alignment = LEFT
        r += 1
    for row in heights:
        ws.cell(row=r, column=1, value="Height implied")
        ws.cell(row=r, column=2, value=f"{row['height_m']:g} m")
        c = ws.cell(row=r, column=3,
                    value=f"{row['top']} over {row['bottom']} — a pairing the "
                          f"section confirms, not the extractor")
        c.alignment = LEFT
        r += 1

    _widths(ws, {1: 6, 2: 46, 3: 8, 4: 12, 5: 46, 6: 7, 7: 16, 8: 52})
    ws.freeze_panes = f"A{head + 1}"
    wb.save(path)
    return path
