#!/usr/bin/env python
"""Engineering drawing PDF -> multi-sheet BOQ workbook, headless.

    venv/bin/python run_pipeline.py                    # the PDF in the project root
    venv/bin/python run_pipeline.py drawing.pdf
    venv/bin/python run_pipeline.py --batch drawings/  # a whole set + roll-up

What happens
------------
    1. Locate the sheet's text from vector geometry      (no model calls)
    2. Title block                                        (1 call)
    3. Transcribe the located text, in montages           (~4 calls)
    4. Grade slab extent from the foundation plan         (1 call)
    5. Grammars turn the transcript into elements         (no model calls)
    6. Optional: derive excavation / blinding / liner / … (no model calls)
    7. core.bom_builder -> core.excel_writer              (Phase 1, unchanged)

Alongside the workbook it writes what a civil team needs to check the takeoff:
a **check print** (the drawing with every extracted value boxed and quoted by
grid square) and a **Verification** sheet with blank sign-off columns.

Draft status: every workbook opens with a red UNVERIFIED DRAFT block listing each
extracted value and its provenance. The reviewed path is `./bin/app.sh`.
"""
from __future__ import annotations

import argparse
import json
import sys
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from core.bom_builder import build_bom
from core.derivation import (
    apply_derived, default_rules, derive, rules_from_findings,
)
from core.excel_writer import write_workbook
from core.filename import build_output_path, sanitise
from core import set_workbook as SW
from core.models import Project
from extractors import qwen_vision as QV
from extractors import rag_examples
from extractors import sheet_grid as SG
from extractors import verification as V
from extractors import workbook_extras as WE
from extractors.models import ExtractionResult
from extractors.ollama_client import OllamaClient
from extractors.pdf_to_image import open_page
from extractors.qwen_vision import ExtractionConfig, PROFILES

ROOT = Path(__file__).resolve().parent
OUTPUT_DIR = ROOT / "output"

# Re-exported so pages/0_Extract.py keeps working against a stable name.
append_audit_sheet = WE.append_audit_sheet


def say(msg: str = "") -> None:
    print(msg, flush=True)


def rule(title: str) -> None:
    say(f"\n{'─' * 74}\n{title}\n{'─' * 74}")


@dataclass
class RunOutcome:
    pdf: Path
    ok: bool
    project: Project | None = None
    result: ExtractionResult | None = None
    workbook: Path | None = None
    check_print: Path | None = None
    error: str = ""
    # Which tags in this drawing are assumptions rather than readings, so the
    # consolidated workbook can shade them.
    derived_tags: set = field(default_factory=set)
    placeholder_tags: set = field(default_factory=set)
    notes: list = field(default_factory=list)


# ---------- PDF discovery ----------
def find_root_pdf() -> Path:
    pdfs = sorted(p for p in ROOT.glob("*.pdf") if not p.name.startswith("."))
    if not pdfs:
        raise SystemExit(f"No PDF in {ROOT}. Pass one: run_pipeline.py drawing.pdf")
    if len(pdfs) > 1:
        names = "\n  ".join(p.name for p in pdfs)
        raise SystemExit(f"Several PDFs in the project root — name one, or use "
                         f"--batch:\n  {names}")
    return pdfs[0]


def find_batch_pdfs(folder: Path) -> list[Path]:
    if not folder.is_dir():
        raise SystemExit(f"--batch expects a folder, got {folder}")
    pdfs = sorted(p for p in folder.rglob("*.pdf") if not p.name.startswith("."))
    if not pdfs:
        raise SystemExit(f"No PDFs under {folder}")
    return pdfs


# ---------- Reporting ----------
def report_extraction(result: ExtractionResult) -> None:
    tb = result.title_block
    rule("Title block")
    for label, value in (("Drawing no", tb.drawing_no), ("Revision", tb.revision),
                         ("Project", tb.project_name),
                         ("Date", f"{tb.date}  (read as {tb.date_raw!r})"
                                  if tb.date else tb.date_raw),
                         ("Drawn by", tb.prepared_by)):
        say(f"  {label:<12} {value or '(not read)'}")

    rule("Pedestal callouts")
    if not result.pedestals:
        say("  (none found)")
    else:
        say(f"  {'Tag':<5} {'L (mm)':>8} {'W (mm)':>8} {'Nos':>5} "
            f"{'Grid':<8} {'Conf':<7} Callout")
        for p in result.pedestals:
            say(f"  {p.tag:<5} {p.length_mm:>8.0f} {p.width_mm:>8.0f} "
                f"{p.quantity:>5} {p.grid_ref or '-':<8} {p.confidence:<7} "
                f"{p.raw_text or '-'}")

    rule("Grade slab")
    if not result.grade_slabs:
        say("  (not read)")
    for s in result.grade_slabs:
        say(f"  {s.tag}: {s.length_mm:g} x {s.width_mm:g} x {s.thickness_mm:g} mm"
            + (f"   TOC EL {s.toc_level}" if s.toc_level else ""))

    if result.findings:
        rule("Also read (review only — not added to the BOQ)")
        labels = {"curb_walls": "Curb wall", "sumps": "Sump", "levels": "Level",
                  "insert_plates": "Insert plate", "epoxy": "Epoxy",
                  "rebar": "Rebar", "thicknesses": "Thickness"}
        for key, label in labels.items():
            for item in result.findings.get(key, [])[:8]:
                detail = ", ".join(f"{k}={v}" for k, v in item.items()
                                   if k not in ("raw_text", "grid_ref",
                                                "source_rect") and v is not None)
                say(f"  {label:<13} {item.get('grid_ref', ''):<8} {detail:<38} "
                    f"\"{item.get('raw_text', '')}\"")

    rule("Model calls")
    for r in result.responses:
        state = "SKIP" if r.skipped else (f"ERROR: {r.error}" if r.error
                                          else f"{r.elapsed_s:>6.1f} s")
        say(f"  {r.stage:<12} {r.region:<18} {r.px_width:>5}x{r.px_height:<5} {state}")
    say(f"  TOTAL {result.total_elapsed_s:.0f} s")

    rule("Notes")
    for n in result.confidence_notes:
        say(f"  • {n}")


# ---------- Config ----------
def build_config(args) -> ExtractionConfig:
    cfg = ExtractionConfig(**PROFILES[args.profile].__dict__)
    if args.tiles:
        try:
            cols, rows = (int(v) for v in args.tiles.lower().split("x"))
        except ValueError:
            raise SystemExit(f"--tiles expects COLSxROWS, got {args.tiles!r}")
        cfg.tile_cols, cfg.tile_rows = cols, rows
    if args.tile_px:
        cfg.tile_px = args.tile_px
    if args.tile_mp:
        cfg.tile_max_pixels = int(args.tile_mp * 1_000_000)
    if args.no_slab:
        cfg.extract_grade_slab = False
    if args.no_rag:
        cfg.use_rag = False
    return cfg


def build_rules(args, result: ExtractionResult):
    rules = rules_from_findings(default_rules(), result.findings)
    wanted = {r.strip() for r in (args.derive or "").split(",") if r.strip()}
    if "all" in wanted:
        for r in rules:
            r.enabled = True
    else:
        for r in rules:
            r.enabled = r.key in wanted
    return rules


# ---------- One drawing ----------
def process_one(pdf_path: Path, args, client: OllamaClient | None) -> RunOutcome:
    say(f"\n{'=' * 74}\n  {pdf_path.name}\n{'=' * 74}")
    try:
        if args.from_json:
            result = ExtractionResult(**json.loads(Path(args.from_json).read_text()))
            say(f"  replayed from {args.from_json} (no model calls)")
        else:
            cfg = build_config(args)
            if cfg.strategy == "montage":
                say("  Strategy: locate text from vector geometry, then transcribe")
            else:
                say(f"  Strategy: blind tile sweep "
                    f"{cfg.tile_cols}x{cfg.tile_rows}")
            result = QV.extract_from_pdf(
                pdf_path, page_number=args.page - 1, config=cfg, client=client,
                progress=lambda label, i, n: say(f"  [{i}/{n}] {label} …"),
                debug_dir=args.debug_dir)

        if args.prepared_by:
            result.title_block.prepared_by = args.prepared_by
        report_extraction(result)

        OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        json_path = OUTPUT_DIR / f"{pdf_path.stem}_extraction.json"
        json_path.write_text(json.dumps(result.model_dump(), indent=2),
                             encoding="utf-8")
        say(f"\n  Extraction JSON → {json_path}")

        project = Project(project_name="", drawing_no="")
        plan = QV.plan_merge(project, result)
        QV.merge_into_project(project, result)

        derived = []
        derived_tags: set[str] = set()
        if args.derive:
            rules = build_rules(args, result)
            derived = derive(project, rules)
            derived_tags = {getattr(i.element, "tag", "") for i in derived}
            added = apply_derived(project, derived)
            if derived:
                rule("Derived quantities (not read from the drawing)")
                for item in derived:
                    say(f"  {item.target_field:<16} "
                        f"{getattr(item.element, 'tag', ''):<12} {item.explanation}")
                say(f"  → added {len(added)} element(s)")

        if args.extract_only:
            say("\n  --extract-only: stopping before the workbook.")
            return RunOutcome(pdf_path, True, project, result)

        if not project.drawing_no:
            project.drawing_no = sanitise(pdf_path.stem)
            say(f"\n  No drawing number read — using the filename "
                f"({project.drawing_no}).")

        discovered = (result.discovery or {}).get("items") or []
        if not (project.pedestals or project.grade_slabs or discovered):
            # Not a failure. A sections sheet has no quantities of its own, and
            # a plan sheet marks positions while referencing sizes elsewhere —
            # both still belong in the set, with whatever they do carry.
            note = ("no BOQ quantities on this sheet"
                    + (f"; position marks counted: "
                       f"{', '.join(f'{k}x{v}' for k, v in result.position_marks.items())}"
                       if result.position_marks else ""))
            say(f"\n  {note}.")
            return RunOutcome(pdf_path, True, project, result, None, None,
                              notes=[note])
        if not (project.pedestals or project.grade_slabs):
            # No element the template knows how to price, but the sheet does
            # specify things. The workbook is still written: its standard tabs
            # come out empty, which is the truth about this drawing, and the
            # Drawing_Items tab carries what it does say.
            say(f"\n  No pedestal or slab on this sheet — writing the workbook "
                f"for the {len(discovered)} item(s) read from the drawing itself.")

        rule("BOM + workbook")
        project.created_at = datetime.now().strftime("%Y-%m-%d %H:%M")
        bom = build_bom(project)
        bom.warnings = QV.extraction_warnings(result, plan) + bom.warnings
        out_path = (Path(args.out) if args.out and not args.batch
                    else build_output_path(project.drawing_no, OUTPUT_DIR))
        written = write_workbook(project, bom, out_path)

        grid = _grid_for(pdf_path, args)
        if not args.no_verify_sheet:
            WE.append_verification_sheet(written, result, grid)
        if derived:
            WE.append_derivation_sheet(written, derived)
        # Not conditional on `derived`: a sheet with nothing to derive from
        # is exactly the sheet whose only content is what it names itself.
        WE.append_discovery_sheet(written, result)
        if args.audit_sheet:
            WE.append_audit_sheet(written, result, plan)

        check = None
        if not args.no_check_print:
            doc, page = open_page(pdf_path, args.page - 1)
            try:
                check = V.build_check_print(
                    page, result, grid=grid,
                    out_path=OUTPUT_DIR / f"{sanitise(project.drawing_no)}_check.png")
            finally:
                doc.close()

        from openpyxl import load_workbook
        say(f"  Sheets  : {', '.join(load_workbook(written).sheetnames)}")
        say(f"  BOM     : {len(bom.lines)} line(s)")
        say(f"  Workbook → {written}")
        if check:
            say(f"  Check print → {check}")

        if args.save_example and result.pedestals:
            ex = rag_examples.save_example(
                drawing_no=project.drawing_no, revision=project.revision,
                image_hash=result.image_hash,
                verified_extraction={"pedestals": [
                    {"tag": p.tag, "length_mm": p.length_mm,
                     "width_mm": p.width_mm, "quantity": p.quantity}
                    for p in result.pedestals]},
                verified_by=args.prepared_by or "", source_pdf=pdf_path.name)
            say(f"  Example → {ex}")

        placeholder_tags = {p.tag for p in project.pedestals
                            if abs(p.height_m - QV.PLACEHOLDER_HEIGHT_M) < 1e-9}
        notes = []
        if result.used_text_layer:
            notes.append("read from the sheet's own text layer (no model calls)")
        if result.position_marks:
            notes.append("position marks counted: " + ", ".join(
                f"{k}x{v}" for k, v in result.position_marks.items()))
        if placeholder_tags:
            notes.append(f"placeholder height on {', '.join(sorted(placeholder_tags))}")
        return RunOutcome(pdf_path, True, project, result, written, check,
                          derived_tags=derived_tags,
                          placeholder_tags=placeholder_tags, notes=notes)

    except Exception as exc:                    # noqa: BLE001 — batch must continue
        return RunOutcome(pdf_path, False, error=f"{type(exc).__name__}: {exc}")


def _grid_for(pdf_path: Path, args):
    if args.no_grid_detect:
        return SG.DEFAULT_GRID
    doc, page = open_page(pdf_path, args.page - 1)
    try:
        return SG.detect_grid(page)
    finally:
        doc.close()


# ---------- Roll-up across a set ----------
def write_rollup(outcomes: list[RunOutcome], path: Path) -> Path:
    """One workbook summing the set, with a per-drawing breakdown.

    A drawing set is priced as a set, but every quantity still has to be
    traceable to the sheet it came from — so the roll-up carries both.
    """
    from collections import defaultdict

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    good = [o for o in outcomes if o.ok and o.project]
    wb = Workbook()
    ws = wb.active
    ws.title = "Set Summary"
    head_fill = PatternFill("solid", fgColor="1F4E78")
    head_font = Font(bold=True, color="FFFFFF")

    ws["A1"] = "DRAWING SET — COMBINED BOQ (UNVERIFIED DRAFT)"
    ws["A1"].font = Font(bold=True, size=14, color="C00000")
    ws["A2"] = (f"{len(good)} drawing(s) processed, "
                f"{len(outcomes) - len(good)} failed   ·   "
                f"{datetime.now().strftime('%Y-%m-%d %H:%M')}")
    ws.merge_cells("A2:F2")

    totals: dict[tuple[str, str, str], float] = defaultdict(float)
    per_drawing: dict[str, dict[tuple[str, str, str], float]] = {}
    for o in good:
        bom = build_bom(o.project)
        key_map: dict[tuple[str, str, str], float] = defaultdict(float)
        for line in bom.lines:
            if line.category == "ROLLUP":
                continue
            k = (line.category, line.item, line.unit)
            totals[k] += line.qty_gross
            key_map[k] += line.qty_gross
        per_drawing[o.project.drawing_no or o.pdf.stem] = key_map

    row = 4
    for col, h in enumerate(["Category", "Item", "Unit", "Qty (gross)",
                             "Drawings"], start=1):
        c = ws.cell(row=row, column=col, value=h)
        c.fill, c.font = head_fill, head_font
    row += 1
    for (cat, item, unit), qty in sorted(totals.items()):
        n = sum(1 for m in per_drawing.values() if (cat, item, unit) in m)
        ws.cell(row=row, column=1, value=cat)
        ws.cell(row=row, column=2, value=item)
        ws.cell(row=row, column=3, value=unit)
        ws.cell(row=row, column=4, value=round(qty, 3)).number_format = "#,##0.000"
        ws.cell(row=row, column=5, value=n)
        row += 1
    for col, w in {1: 22, 2: 46, 3: 8, 4: 14, 5: 10}.items():
        ws.column_dimensions[chr(64 + col)].width = w

    ws2 = wb.create_sheet("Per Drawing")
    for col, h in enumerate(["Drawing", "Status", "Pedestals", "Slabs",
                             "BOQ lines", "Workbook", "Error"], start=1):
        c = ws2.cell(row=1, column=col, value=h)
        c.fill, c.font = head_fill, head_font
    for i, o in enumerate(outcomes, start=2):
        ws2.cell(row=i, column=1, value=o.pdf.name)
        ws2.cell(row=i, column=2, value="ok" if o.ok else "FAILED")
        if o.project:
            ws2.cell(row=i, column=3, value=len(o.project.pedestals))
            ws2.cell(row=i, column=4, value=len(o.project.grade_slabs))
            ws2.cell(row=i, column=5, value=len(build_bom(o.project).lines))
        ws2.cell(row=i, column=6, value=o.workbook.name if o.workbook else "")
        c = ws2.cell(row=i, column=7, value=o.error)
        c.alignment = Alignment(wrap_text=True)
    for col, w in {1: 42, 2: 10, 3: 11, 4: 8, 5: 11, 6: 40, 7: 50}.items():
        ws2.column_dimensions[chr(64 + col)].width = w

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# ---------- CLI ----------
def _reads_from_text(pdf: Path, page: int = 1) -> bool:
    """Whether this sheet can be read without a single model call.

    `page` is 1-based, matching --page on the command line.
    """
    try:
        import fitz

        from extractors import text_layer as TL
        with fitz.open(pdf) as doc:
            return TL.has_usable_text(doc[page - 1])
    except Exception:                                    # noqa: BLE001
        return False                                     # assume it needs vision


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(
        description="Engineering drawing PDF -> BOQ workbook via a local "
                    "Qwen2.5-VL model.",
        formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pdf", nargs="?", help="drawing PDF (default: the project root one)")
    ap.add_argument("--batch", help="folder of PDFs; one workbook each plus a roll-up")
    ap.add_argument("--page", type=int, default=1, help="1-based page number")
    ap.add_argument("--profile", choices=sorted(PROFILES), default="thorough",
                    help="thorough/quick = montage (fast, default); "
                         "sweep/fast = blind tile sweep for scanned sheets")
    ap.add_argument("--derive", metavar="RULES",
                    help="derive dependent quantities: 'all', or a comma list of "
                         "excavation,fill,blinding,liner,epoxy,curb,joints")
    ap.add_argument("--out", help="output .xlsx path (single drawing only)")
    ap.add_argument("--audit-sheet", action="store_true",
                    help="append the Extraction_Log sheet")
    ap.add_argument("--no-verify-sheet", action="store_true",
                    help="skip the Verification sign-off sheet")
    ap.add_argument("--no-check-print", action="store_true",
                    help="skip the annotated check print PNG")
    ap.add_argument("--no-grid-detect", action="store_true",
                    help="assume a 16x16 border grid instead of measuring it")
    ap.add_argument("--extract-only", action="store_true")
    ap.add_argument("--from-json", help="rebuild from a saved extraction JSON")
    ap.add_argument("--tiles", help="sweep grid override, e.g. 6x4")
    ap.add_argument("--tile-px", type=int)
    ap.add_argument("--tile-mp", type=float, help="pixel budget per image, MP")
    ap.add_argument("--no-slab", action="store_true")
    ap.add_argument("--no-rag", action="store_true")
    ap.add_argument("--debug-dir", help="save every rendered crop here")
    ap.add_argument("--prepared-by", default="")
    ap.add_argument("--save-example", action="store_true")
    ap.add_argument("--model", help="override the Ollama model tag")
    args = ap.parse_args(argv)

    if args.batch:
        pdfs = find_batch_pdfs(Path(args.batch))
    else:
        pdfs = [Path(args.pdf) if args.pdf else find_root_pdf()]
        if not pdfs[0].exists():
            raise SystemExit(f"PDF not found: {pdfs[0]}")

    client = None
    if not args.from_json:
        client = OllamaClient(model=args.model) if args.model else OllamaClient()
        ok, msg = client.health()
        say(f"Model: {msg}")
        if not ok:
            # A sheet that kept its text layer costs no model calls at all —
            # refusing to run it because Ollama is down would be refusing work
            # the machine can do in a fifth of a second. Only stop when at least
            # one drawing in the run actually needs the model.
            needs_model = [p for p in pdfs if not _reads_from_text(p, args.page)]
            if needs_model:
                names = ", ".join(p.name for p in needs_model[:3])
                more = f" and {len(needs_model) - 3} more" if len(needs_model) > 3 else ""
                raise SystemExit(
                    f"{len(needs_model)} drawing(s) have no text layer and need "
                    f"the vision model ({names}{more}).\n"
                    "Start the model host first:\n  ollama serve\n"
                    "  ollama pull qwen2.5vl:7b")
            say("Every drawing in this run has its own text layer — "
                "continuing without the model.")

    if len(pdfs) > 1:
        say(f"Batch: {len(pdfs)} drawing(s) from {args.batch}")

    outcomes = [process_one(p, args, client) for p in pdfs]

    failed = [o for o in outcomes if not o.ok]
    if len(pdfs) > 1:
        rule("Batch complete")
        for o in outcomes:
            say(f"  {'ok  ' if o.ok else 'FAIL'} {o.pdf.name}"
                + (f" — {o.error}" if o.error else ""))
        entries = [
            SW.build_entry(o.project.drawing_no, o.project,
                           source_pdf=str(o.pdf), derived_tags=o.derived_tags,
                           placeholder_tags=o.placeholder_tags, notes=o.notes,
                           position_marks=(o.result.position_marks
                                           if o.result else {}),
                           discovery=(o.result.discovery if o.result else {}))
            for o in outcomes if o.ok and o.project]
        if entries:
            consolidated = SW.write_set_workbook(
                entries, OUTPUT_DIR / "SET_BOQ.xlsx",
                project_name=next((o.project.project_name for o in outcomes
                                   if o.ok and o.project and o.project.project_name),
                                  ""))
            say(f"\n  Consolidated BOQ → {consolidated}")
        rollup = write_rollup(outcomes, OUTPUT_DIR / "SET_ROLLUP.xlsx")
        say(f"  Status roll-up  → {rollup}")

    rule("Done — UNVERIFIED DRAFT")
    say("  Every workbook needs an engineer's check before pricing:")
    say("    • enter real pedestal heights (not printed on these callouts)")
    say("    • mark up the Verification sheet against the check print")
    say("    • add anything the drawing does not dimension (sump depth, rebar)")
    return 1 if failed and len(pdfs) == 1 else 0


if __name__ == "__main__":
    sys.exit(main())
