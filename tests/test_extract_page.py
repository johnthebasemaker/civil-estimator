"""Execution tests for the Drawing → BOQ page.

Streamlit pages fail at *render* time, not import time: a widget argument that
does not exist in the installed version, or a None slipping into a data editor,
looks fine to a linter and dies when someone opens the tab. That is exactly how
`st.image(use_container_width=...)` shipped broken.

`AppTest` runs the page script in-process, so these are real executions with no
browser and no model.
"""
from __future__ import annotations

from pathlib import Path

import pytest

from core.models import Project
from extractors.models import (
    ExtractionResult, GradeSlabExtraction, PedestalExtraction,
    RegionResponse, TitleBlockExtraction,
)

AppTest = pytest.importorskip("streamlit.testing.v1").AppTest

ROOT = Path(__file__).resolve().parent.parent
from ui.auth import SESSION_KEY as AUTH_KEY

PAGE = str(ROOT / "pages" / "0_Extract.py")
from tests.sample_drawing import sample_pdf

SAMPLE_PDF = sample_pdf()

TRUTH = {"P1": (600, 500, 2), "P2": (500, 500, 20), "P3": (350, 350, 11),
         "P6": (450, 450, 6), "P7": (500, 500, 4)}


def _extraction() -> ExtractionResult:
    r = ExtractionResult(
        source_pdf=str(SAMPLE_PDF), model="qwen2.5vl:7b", profile="thorough",
        image_hash="5110010820424080", total_elapsed_s=250.0,
        montages_sent=4, text_blocks_found=171,
        title_block=TitleBlockExtraction(
            drawing_no="MD-522-8110-EG-CV-LAD-0107", revision="C01",
            project_name="MAADEN PHOSPHATE 3 PHASE 1", date="2025-09-30",
            prepared_by="P.PARDULE", confidence="high"),
        pedestals=[PedestalExtraction(
            tag=t, length_mm=l, width_mm=w, quantity=q, regex_validated=True,
            raw_text=f"TYP DETAIL OF PEDESTAL {t}({l}x{w}) {q}Nos",
            confidence="medium", seen_in_regions=["transcript"])
            for t, (l, w, q) in TRUTH.items()],
        grade_slabs=[GradeSlabExtraction(
            tag="GS-01", length_mm=24430, width_mm=8600, thickness_mm=300,
            toc_level="97.550", confidence="medium")],
        responses=[RegionResponse(region="montage_1", stage="transcribe",
                                  px_width=1200, px_height=631, elapsed_s=44.0)],
        transcribed_lines=["TYP DETAIL OF PEDESTAL", "P1(600x500) 2Nos",
                           "CURB WALL 150THK X 150HIGH", "TOC EL 97.850"],
        findings={"curb_walls": [{"thickness_mm": 150.0, "height_mm": 150.0,
                                  "raw_text": "CURB WALL 150THK X 150HIGH"}],
                  "levels": [{"kind": "TOC", "level_m": 97.85,
                              "raw_text": "TOC EL 97.850"}],
                  "sumps": [], "insert_plates": [], "epoxy": [], "rebar": [],
                  "thicknesses": []},
        confidence_notes=["located 171 text blocks from vector geometry"],
    )
    return r


def _open(at, name: str) -> bool:
    """Press a drawing's Open button, the way a person opens one."""
    button = next((b for b in at.button
                   if b.key and b.key.startswith("open::") and b.key.endswith(name)),
                  None)
    if button is None:
        return False
    button.click().run()
    return True


def _titles(at) -> list[str]:
    """Headers and subheaders: the workspace titles its sections with both."""
    return [h.value for h in at.header] + [s.value for s in at.subheader]


def _page_with_drawing(extraction=None, project=None):
    """Drive the page the way a person does: open a drawing, then read it.

    Order matters. The page clears any stale extraction when the open drawing
    changes, so seeding `extraction` before the drawing is opened would be wiped
    — exactly as it should be.
    """
    at = AppTest.from_file(PAGE, default_timeout=300)
    at.session_state[AUTH_KEY] = True          # the page is behind a login gate
    if project is not None:
        at.session_state["project"] = project
    at.run()
    _open(at, SAMPLE_PDF.name)

    if extraction is not None:
        at.session_state["extraction"] = extraction
        at.run()
    return at


@pytest.fixture(autouse=True)
def _sandboxed_queue(tmp_path, monkeypatch):
    """A queue, a cache and a drawing library of this test's own.

    Without this the page renders against whatever is in the developer's real
    `data/jobs.db` and uploads folder, so the same test passes or fails
    depending on what someone extracted — or deleted — this morning. That is
    exactly how this file went red once the uploads folder was emptied.

    The library holds symlinks to the real drawings: the page sees them as
    ordinary files, and nothing is copied.
    """
    monkeypatch.setenv("CIVIL_ESTIMATOR_JOBS_DB", str(tmp_path / "jobs.db"))
    monkeypatch.setenv("CIVIL_ESTIMATOR_CACHE_DIR", str(tmp_path / "cache"))
    drawings = tmp_path / "Drawings"
    drawings.mkdir()
    for pdf in (SAMPLE_PDF, ROOT / "Drawings" / "MD-522-8110-EG-CV-LAD-0101_C01.pdf"):
        if pdf.exists():
            (drawings / pdf.name).symlink_to(pdf.resolve())
    monkeypatch.setenv("CIVIL_ESTIMATOR_DRAWING_DIRS", str(drawings))
    monkeypatch.setenv("CIVIL_ESTIMATOR_UPLOAD_DIR", str(tmp_path / "uploads"))
    # Workbooks, check prints and SET_BOQ.xlsx go to the test's own folder.
    # Before this, every run wrote them into the real output/ — overwriting
    # the SET_BOQ.xlsx someone had just built.
    from ui.workspace import common as _C
    monkeypatch.setattr(_C, "OUTPUT_DIR", tmp_path / "output")
    monkeypatch.setattr(_C, "PROJECT_DIR", tmp_path / "output" / "projects")
    return tmp_path


def _load_worker():
    """bin/ is a script directory, not a package, so load the worker by path."""
    import importlib.util
    spec = importlib.util.spec_from_file_location("worker", ROOT / "bin" / "worker.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _run_the_worker_once(tmp_path, monkeypatch):
    """What `bin/worker.py` does for one queued job, in-process.

    Its human-readable JSON copy is redirected into the test's folder, so a test
    run never overwrites the real `output/<drawing>_extraction.json`.
    """
    from core import jobstore as JS

    W = _load_worker()
    monkeypatch.setattr(W, "OUTPUT_DIR", tmp_path / "output")
    job = JS.claim_next("test-worker")
    assert job is not None, "nothing was queued"
    return W.process(job, None), job


@pytest.mark.skipif(not SAMPLE_PDF.exists(), reason="sample drawing not in repo root")
class TestExtractPageRenders:
    def test_opens_on_the_drawing_list(self):
        at = AppTest.from_file(PAGE, default_timeout=120)
        at.session_state[AUTH_KEY] = True
        at.run()
        assert not at.exception
        assert any(b.label == "Open" for b in at.button)
        assert "1 · Sheet" not in _titles(at), "nothing is open until Open is pressed"

    def test_page_is_behind_the_login_gate(self):
        """Every page carries its own gate: Streamlit runs each as a separate
        script, so a gate on the home page alone would be bypassed by
        navigating straight here."""
        at = AppTest.from_file(PAGE, default_timeout=120).run()
        assert not at.exception
        assert not at.tabs, "the workspace must not render behind the gate"
        assert any(t.label == "Password" for t in at.text_input)

    def test_renders_sheet_metrics_without_calling_the_model(self):
        at = _page_with_drawing()
        assert not at.exception
        labels = {m.label: m.value for m in at.metric}
        assert labels["Sheet"] == "A0"
        assert int(labels["Text blocks found"]) > 0
        assert int(labels["Model calls needed"]) < 10

    def test_review_and_generate_render_once_extraction_exists(self):
        at = _page_with_drawing(_extraction(), Project(project_name="", drawing_no=""))
        assert not at.exception
        titles = _titles(at)
        assert "3 · Review and correct" in titles
        assert "4 · Generate this drawing's BOQ" in titles
        assert any("Generate BOQ Excel" in b.label for b in at.button)

    def test_pedestal_grid_is_prefilled_and_flags_placeholder_height(self):
        at = _page_with_drawing(_extraction(), Project(project_name="", drawing_no=""))
        assert not at.exception
        assert any("height" in w.value.lower() for w in at.warning)
        assert any("placeholder" in w.value.lower() for w in at.warning)

    def test_extra_findings_are_shown_but_not_in_the_project(self):
        at = _page_with_drawing(_extraction(), Project(project_name="", drawing_no=""))
        assert not at.exception
        assert at.session_state.project.curb_walls == []

    def test_generate_button_writes_a_workbook(self):
        at = _page_with_drawing(_extraction(), Project(project_name="", drawing_no=""))
        btn = next(b for b in at.button if "Generate BOQ Excel" in b.label)
        btn.click().run()
        assert not at.exception, [e.value for e in at.exception]
        written = at.session_state["last_workbook"]
        assert written and Path(written).exists()

        from openpyxl import load_workbook
        wb = load_workbook(written)
        # The mandated ten-sheet template, in order. Gaps_and_Assumptions is
        # inserted straight after the Summary when the drawing left something
        # unstated, which is almost always — so it is skipped here rather than
        # allowed to shift the ten.
        template = [n for n in wb.sheetnames if n != "Gaps_and_Assumptions"]
        expected = ["Summary", "Assumptions", "Pedestals", "Slab", "Sump",
                    "Joints", "Coating", "Embedments", "Rebar_BBS", "Costing"]
        assert template[:10] == expected
        assert wb.sheetnames[1] == "Gaps_and_Assumptions", \
            "the gaps belong where they will be read, not at the end"
        rows = [r for r in wb["Pedestals"].iter_rows(values_only=True)
                if r and r[0] in TRUTH]
        assert len(rows) == len(TRUTH)

    def test_generated_project_carries_the_extracted_metadata(self):
        at = _page_with_drawing(_extraction(), Project(project_name="", drawing_no=""))
        next(b for b in at.button if "Generate BOQ Excel" in b.label).click().run()
        p = at.session_state.project
        assert p.drawing_no == "MD-522-8110-EG-CV-LAD-0107"
        assert {x.tag for x in p.pedestals} == set(TRUTH)
        assert p.grade_slabs[0].length_m == 24.430


class TestStreamlitCompat:
    def test_image_flag_matches_installed_streamlit(self):
        import streamlit as st
        from extractors.st_compat import _IMAGE_FULL_WIDTH
        import inspect
        params = inspect.signature(st.image).parameters
        assert _IMAGE_FULL_WIDTH in params, (
            f"st.image on streamlit {st.__version__} has no {_IMAGE_FULL_WIDTH}")


@pytest.mark.skipif(not SAMPLE_PDF.exists(), reason="sample drawing not in repo root")
class TestDiscoveryPanel:
    """Items the drawing names itself must reach the reviewer, not just the file.

    The five sheets that carry no pedestal and no slab have nothing else to show
    on this page — if the panel does not render, the person checking the takeoff
    sees an empty review grid and concludes the extractor failed.
    """

    def _extraction_with_discovery(self):
        r = _extraction()
        r.discovery = {
            "items": [
                {"kind": "thickness",
                 "description": "Acid Resistant Epoxy Coating, 4 mm thick",
                 "spec": {"thickness_mm": 4.0}, "uom": "m2", "qty": None,
                 "basis": "4 mm thick — area not stated on this sheet",
                 "occurrences": 3, "confirm": True,
                 "source": "4MM THK ACID RESISTANT", "grid_ref": "J-15 → P-14"},
                {"kind": "rebar", "description": "D25 bars — 16 Nos per element",
                 "spec": {"count": 16.0, "diameter_mm": 25.0}, "uom": "Nos",
                 "qty": 16.0, "basis": "16-D25 — bars per element",
                 "occurrences": 5, "confirm": False, "source": "16-D25",
                 "grid_ref": "I-9 → P-7"},
            ],
            "levels": [], "grades": [],
            "heights": [{"height_m": 2.88, "top": "BOBP EL. 98.380",
                         "bottom": "BOC EL. 95.500"}],
            "specifications": [{"value": 35.0, "unit": "MPa",
                                "text": "MINIMUM STRENGTH OF 35 MPa"}],
            "measured": 1, "to_confirm": 1,
        }
        return r

    def test_the_page_renders_with_discovered_items(self):
        at = _page_with_drawing(extraction=self._extraction_with_discovery())
        assert not at.exception

    def test_the_panel_is_shown_and_open(self):
        at = _page_with_drawing(extraction=self._extraction_with_discovery())
        panel = next((e for e in at.expander
                      if "Read from this drawing" in e.label), None)
        assert panel is not None, [e.label for e in at.expander]
        # AppTest does not expose the open/closed flag, so check the content
        # instead: the table of items has to be inside this panel.
        assert panel.dataframe, "the panel rendered without its item table"

    def test_the_panel_counts_what_carries_a_quantity(self):
        at = _page_with_drawing(extraction=self._extraction_with_discovery())
        label = next(e.label for e in at.expander
                     if "Read from this drawing" in e.label)
        assert "2 item(s)" in label and "1 with a stated quantity" in label

    def test_no_panel_when_the_sheet_specified_nothing(self):
        at = _page_with_drawing(extraction=_extraction())
        assert not any("Read from this drawing" in e.label for e in at.expander)


@pytest.mark.skipif(not SAMPLE_PDF.exists(), reason="sample drawing not in repo root")
class TestAMissingFigureNeverBlocks:
    """The button used to be disabled unless the drawing produced a pedestal or
    a slab. Five sheets in this set produce neither, so five sheets produced
    nothing — and the missing figure was still missing afterwards."""

    def _bare(self):
        """An extraction with no pedestal and no slab, like a sections sheet."""
        r = _extraction()
        r.pedestals, r.grade_slabs = [], []
        r.discovery = {
            "items": [{"kind": "thickness",
                       "description": "Acid Resistant Epoxy Coating, 4 mm thick",
                       "spec": {"thickness_mm": 4.0}, "uom": "m2", "qty": None,
                       "basis": "4 mm thick — area not stated on this sheet",
                       "occurrences": 3, "confirm": True,
                       "source": "4MM THK ACID RESISTANT", "grid_ref": "J-15"}],
            "levels": [], "heights": [], "grades": [], "specifications": [],
            "measured": 0, "to_confirm": 1}
        return r

    def test_the_generate_button_is_offered(self):
        at = _page_with_drawing(extraction=self._bare())
        assert not at.exception
        gen = next((b for b in at.button
                    if b.label.startswith("🧾") and "Generate BOQ" in b.label), None)
        assert gen is not None, [b.label for b in at.button]
        assert not gen.disabled, "a sections sheet still deserves a workbook"

    def test_the_page_says_why_the_workbook_is_still_worth_having(self):
        at = _page_with_drawing(extraction=self._bare())
        assert any("still worth having" in i.value for i in at.info)

    def test_the_gaps_are_shown_before_the_button(self):
        at = _page_with_drawing(extraction=self._bare())
        assert any("gap(s) and assumption(s)" in e.label for e in at.expander)

    def test_it_generates(self, tmp_path):
        at = _page_with_drawing(extraction=self._bare())
        next(b for b in at.button
             if b.label.startswith("🧾") and "Generate BOQ" in b.label).click().run()
        assert not at.exception
        # AppTest's session state is a SafeSessionState and has no .get().
        assert "last_workbook" in at.session_state
        assert Path(at.session_state["last_workbook"]).exists()

    def test_the_workbook_carries_the_gaps_sheet(self):
        from openpyxl import load_workbook

        at = _page_with_drawing(extraction=self._bare())
        next(b for b in at.button
             if b.label.startswith("🧾") and "Generate BOQ" in b.label).click().run()
        names = load_workbook(at.session_state["last_workbook"]).sheetnames
        assert "Gaps_and_Assumptions" in names
        assert "Drawing_Items" in names


TEXT_LAYER_PDF = ROOT / "Drawings" / "MD-522-8110-EG-CV-LAD-0101_C01.pdf"


@pytest.mark.skipif(not TEXT_LAYER_PDF.exists(), reason="drawings not present")
class TestASheetThatNeedsNoModel:
    """Found by driving the page with Ollama stopped: a drawing that is read
    from its own text layer was still refusing to extract, because the button
    was gated on the model being reachable rather than on the model being
    needed."""

    def _page(self):
        at = AppTest.from_file(PAGE, default_timeout=300)
        at.session_state[AUTH_KEY] = True
        at.run()
        if not _open(at, TEXT_LAYER_PDF.name):
            pytest.skip("drawing not in the library")
        return at

    def test_the_page_says_no_model_calls_are_needed(self):
        at = self._page()
        calls = next(m for m in at.metric if m.label == "Model calls needed")
        assert calls.value == "0"

    def test_it_says_so_in_words_too(self):
        at = self._page()
        assert any("No model calls at all" in s.value for s in at.success)

    def test_the_extract_button_is_not_blocked_by_a_stopped_model(self):
        at = self._page()
        button = next(b for b in at.button if b.label.startswith("🔍"))
        assert not button.disabled

    def test_extracting_works_and_costs_no_model_time(self, tmp_path, monkeypatch):
        """The button queues the drawing; the worker reads it; the page picks
        the result up. The page itself never runs the extractor — see
        tests/test_pages_never_run_the_model.py."""
        at = self._page()
        next(b for b in at.button if b.label.startswith("🔍")).click().run()
        assert not at.exception
        assert "extraction" not in at.session_state, \
            "the page must not read the drawing itself"

        outcome, _ = _run_the_worker_once(tmp_path, monkeypatch)
        assert outcome == "done"
        at.run()
        assert not at.exception
        result = at.session_state["extraction"]
        assert result.used_text_layer
        assert result.montages_sent == 0


@pytest.mark.skipif(not SAMPLE_PDF.exists(), reason="sample drawing not present")
class TestADrawingReadBeforeOpensInstantly:
    """The management case: a drawing somebody already read must come back from
    the saved extraction, with the model switched off, and without queueing."""

    @pytest.fixture
    def saved(self):
        from core import extract_cache as CACHE
        from core import jobstore as JS

        CACHE.store(JS.fingerprint(SAMPLE_PDF, "thorough"), _extraction())

    def test_it_opens_without_the_model(self, saved):
        at = _page_with_drawing()
        assert not at.exception
        assert "extraction" in at.session_state
        assert any("saved reading" in s.value for s in at.success)
        assert "3 · Review and correct" in _titles(at)

    def test_nothing_is_queued_to_open_it(self, saved):
        from core import jobstore as JS

        _page_with_drawing()
        assert JS.list_jobs() == []

    def test_reading_it_again_is_offered_but_not_forced(self, saved):
        at = _page_with_drawing()
        again = next(b for b in at.button if b.label.startswith("🔁"))
        assert "Read again" in again.label
        assert again.proto.type != "primary"


@pytest.mark.skipif(not SAMPLE_PDF.exists(), reason="sample drawing not present")
class TestReadingGoesThroughTheQueue:
    def test_the_button_queues_one_drawing(self):
        from core import jobstore as JS

        at = _page_with_drawing()
        next(b for b in at.button if b.label == "🔍 Read this drawing").click().run()
        assert not at.exception
        jobs = JS.list_jobs()
        assert len(jobs) == 1
        assert jobs[0].drawing_name == SAMPLE_PDF.name
        assert jobs[0].page == 0 and jobs[0].profile == "thorough"

    def test_while_it_waits_the_page_offers_a_stop_and_no_review(self):
        at = _page_with_drawing()
        next(b for b in at.button if b.label == "🔍 Read this drawing").click().run()
        assert any(b.label == "⏹ Stop" for b in at.button)
        assert "3 · Review and correct" not in _titles(at)

    def test_stop_takes_it_off_the_queue(self):
        from core import jobstore as JS

        at = _page_with_drawing()
        next(b for b in at.button if b.label == "🔍 Read this drawing").click().run()
        next(b for b in at.button if b.label == "⏹ Stop").click().run()
        assert not at.exception
        assert JS.list_jobs()[0].state == JS.CANCELLED

    def test_the_finished_reading_is_picked_up(self):
        from core import extract_cache as CACHE
        from core import jobstore as JS

        at = _page_with_drawing()
        next(b for b in at.button if b.label == "🔍 Read this drawing").click().run()
        job = JS.claim_next("test-worker")
        CACHE.store(job.fingerprint, _extraction())
        JS.finish(job.id)
        at.run()
        assert not at.exception
        assert at.session_state["extraction"].title_block.drawing_no == \
            "MD-522-8110-EG-CV-LAD-0107"
        assert any("Read and saved" in s.value for s in at.success)


@pytest.mark.skipif(not SAMPLE_PDF.exists(), reason="sample drawing not present")
class TestAFailureIsSaidPlainly:
    RAW = ("the vision model is unavailable: Cannot reach Ollama at "
           "http://127.0.0.1:11434: <urlopen error [Errno 61] Connection refused>")

    @pytest.fixture
    def failed(self):
        from core import jobstore as JS

        at = _page_with_drawing()
        next(b for b in at.button if b.label == "🔍 Read this drawing").click().run()
        job = JS.claim_next("test-worker")
        JS.fail(job.id, self.RAW)
        at.run()
        return at

    def test_the_headline_is_in_words(self, failed):
        headline = " ".join(e.value for e in failed.error)
        assert "reading service was offline" in headline
        assert "Errno" not in headline and "urlopen" not in headline

    def test_the_raw_message_is_one_click_away(self, failed):
        assert any(e.label == "Technical details" for e in failed.expander)
        assert any(self.RAW in c.value for c in failed.code)

    def test_it_offers_to_try_again(self, failed):
        assert any(b.label == "🔍 Try reading it again" for b in failed.button)


@pytest.mark.skipif(not SAMPLE_PDF.exists(), reason="sample drawing not present")
class TestAReissuedDrawingIsNotShownTheOldReading:
    """Same file name, new bytes: the old result belongs to the old drawing."""

    def test_the_stale_result_is_not_loaded(self, tmp_path, monkeypatch):
        from core import extract_cache as CACHE
        from core import jobstore as JS

        folder = tmp_path / "reissue"
        folder.mkdir()
        pdf = folder / SAMPLE_PDF.name
        pdf.write_bytes(SAMPLE_PDF.read_bytes())
        monkeypatch.setenv("CIVIL_ESTIMATOR_DRAWING_DIRS", str(folder))

        job = JS.enqueue([pdf], owner="shared")[0]
        JS.claim_next("test-worker")
        CACHE.store(job.fingerprint, _extraction())
        JS.finish(job.id)

        with pdf.open("ab") as fh:                 # the reissue
            fh.write(b"\n% revised\n")

        at = _page_with_drawing()
        assert not at.exception
        assert "extraction" not in at.session_state
        assert any(b.label == "🔍 Read this drawing" for b in at.button)
