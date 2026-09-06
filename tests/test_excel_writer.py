"""Smoke tests for excel_writer — verify workbook opens, has all sheets,
and key cells contain expected values/formulas.
"""
from __future__ import annotations
import tempfile
from pathlib import Path
import pytest
from openpyxl import load_workbook
from core.models import (
    Project, Pedestal, GradeSlab, Sump, Joint, Excavation, PCCBlinding,
    EpoxyCoating, Embedment, RebarBar,
)
from core.bom_builder import build_bom
from core.excel_writer import write_workbook


@pytest.fixture
def sample_project():
    return Project(
        project_name="Maaden Phosphate 3", drawing_no="MD-522-8110-EG-CV-LAD-0107",
        revision="C01", prepared_by="Johnson Andrew",
        pedestals=[
            Pedestal(tag="P1", length_m=0.6, width_m=0.5, height_m=0.8, quantity=2),
            Pedestal(tag="P2", length_m=0.5, width_m=0.5, height_m=0.8, quantity=20),
        ],
        grade_slabs=[GradeSlab(tag="GS-01", length_m=25, width_m=13, thickness_m=0.25)],
        sumps=[Sump(tag="SP-01", outer_length_m=2.0, outer_width_m=2.0, depth_m=1.5)],
        joints=[Joint(tag="EJ-01", joint_type="expansion", length_m=25.5, has_waterstop=True)],
        excavations=[Excavation(tag="EXC-01", length_m=27, width_m=15, depth_m=1.5)],
        pcc_blindings=[PCCBlinding(tag="PCC-01", length_m=27, width_m=15)],
        epoxy_coatings=[EpoxyCoating(tag="EP-01", area_m2=325, coats=2)],
        embedments=[Embedment(tag="IP-01", embedment_type="insert_plate",
                              quantity=10, size_description="200x200x10",
                              unit_weight_kg=3.14)],
        rebar_bars=[RebarBar(tag="B1", parent_element="GS-01", diameter_mm=16,
                             cut_length_m=25.0, nos_per_element=40, parent_quantity=1)],
    )


@pytest.fixture
def output_wb(sample_project, tmp_path):
    bom = build_bom(sample_project)
    path = tmp_path / "test_output.xlsx"
    write_workbook(sample_project, bom, path)
    return load_workbook(path, data_only=False)


def test_workbook_has_all_sheets(output_wb):
    expected = {"Summary", "Pedestals", "Slab", "Sump", "Joints",
                "Coating", "Embedments", "Rebar_BBS", "Assumptions", "Costing"}
    assert expected.issubset(set(output_wb.sheetnames))

def test_summary_is_first_sheet(output_wb):
    assert output_wb.sheetnames[0] == "Summary"

def test_assumptions_contains_wastage_values(output_wb):
    ws = output_wb["Assumptions"]
    # Scan column B for wastage values
    b_values = [ws.cell(row=r, column=2).value for r in range(1, 20)]
    assert 3.0 in b_values  # concrete
    assert 5.0 in b_values  # rebar
    assert 10.0 in b_values  # formwork

def test_summary_has_gross_formula(output_wb):
    ws = output_wb["Summary"]
    # Find first non-rollup data row and check column G is a formula
    for r in range(5, 80):
        val = ws.cell(row=r, column=7).value
        if isinstance(val, str) and val.startswith("=E"):
            assert "(1+F" in val  # wastage formula shape
            return
    pytest.fail("No qty_gross formula found in Summary")

def test_costing_has_amount_formula(output_wb):
    ws = output_wb["Costing"]
    for r in range(5, 80):
        val = ws.cell(row=r, column=7).value
        if isinstance(val, str) and val.startswith("=E") and "*F" in val:
            return
    pytest.fail("No amount formula found in Costing")

def test_costing_excludes_rollup(output_wb, sample_project):
    from core.bom_builder import build_bom
    bom = build_bom(sample_project)
    rollup_items = {l.item for l in bom.lines if l.category == "ROLLUP"}
    ws = output_wb["Costing"]
    for r in range(5, 200):
        item = ws.cell(row=r, column=3).value
        if item in rollup_items:
            pytest.fail(f"Rollup item leaked into Costing: {item}")

def test_pedestals_sheet_has_data(output_wb):
    ws = output_wb["Pedestals"]
    # P1 tag should appear in the sheet
    for row in ws.iter_rows(values_only=True):
        if row and "P1" in [str(c) for c in row if c]:
            return
    pytest.fail("P1 not found in Pedestals sheet")

def test_rebar_bbs_has_total_row(output_wb):
    ws = output_wb["Rebar_BBS"]
    found = False
    for r in range(5, 30):
        if ws.cell(row=r, column=1).value == "TOTAL (kg)":
            g_val = ws.cell(row=r, column=7).value
            assert isinstance(g_val, str) and g_val.startswith("=SUM")
            found = True
    assert found

def test_workbook_written_to_disk(sample_project, tmp_path):
    bom = build_bom(sample_project)
    path = tmp_path / "output.xlsx"
    result = write_workbook(sample_project, bom, path)
    assert result.exists()
    assert result.stat().st_size > 3000  # non-trivial workbook

def test_all_sheets_have_drawing_no_and_rev(output_wb, sample_project):
    """Every sheet's title block should include the drawing no and revision."""
    for sheet_name in ["Summary", "Pedestals", "Slab", "Sump", "Joints",
                       "Coating", "Embedments", "Rebar_BBS", "Assumptions", "Costing"]:
        ws = output_wb[sheet_name]
        row2 = ws["A2"].value or ""
        assert sample_project.drawing_no in row2, f"Drawing no missing on {sheet_name}"
        assert sample_project.revision in row2, f"Revision missing on {sheet_name}"

def test_all_sheets_have_prepared_by_and_created(output_wb, sample_project):
    """Every sheet's title block should include 'Prepared by' and 'Created'."""
    for sheet_name in ["Summary", "Pedestals", "Slab", "Sump", "Joints",
                       "Coating", "Embedments", "Rebar_BBS", "Assumptions", "Costing"]:
        ws = output_wb[sheet_name]
        row3 = ws["A3"].value or ""
        assert "Prepared by" in row3, f"Prepared by missing on {sheet_name}"
        assert "Created" in row3, f"Created timestamp missing on {sheet_name}"

def test_assumptions_has_source_drawing_row(output_wb):
    ws = output_wb["Assumptions"]
    found = False
    for r in range(1, 20):
        if ws.cell(row=r, column=1).value == "Source Drawing (PDF)":
            found = True
            break
    assert found, "Source Drawing row missing from Assumptions"

def test_assumptions_has_generated_row(output_wb):
    ws = output_wb["Assumptions"]
    found = False
    for r in range(1, 20):
        if ws.cell(row=r, column=1).value == "Generated":
            found = True
            break
    assert found, "Generated timestamp row missing from Assumptions"
