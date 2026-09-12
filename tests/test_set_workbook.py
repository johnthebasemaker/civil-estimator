"""The consolidated workbook for a drawing set.

Shape asked for: one Summary sheet — Sl. #, Description, UoM, then one quantity
column per drawing — with per-drawing/activity detail sheets behind it, and
assumptions marked rather than blended in with values read off a drawing.
"""
from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from core import set_workbook as SW
from core.models import GradeSlab, Pedestal, Project


def _project(drawing_no: str, ped_qty: int = 2, slab: bool = True) -> Project:
    p = Project(project_name="Set test", drawing_no=drawing_no)
    p.pedestals.append(Pedestal(tag="P1", length_m=0.6, width_m=0.5,
                                height_m=0.8, quantity=ped_qty))
    if slab:
        p.grade_slabs.append(GradeSlab(tag="GS-01", length_m=24.43,
                                       width_m=8.6, thickness_m=0.3))
    return p


def _entries():
    return [
        SW.build_entry("MD-522-8110-EG-CV-LAD-0106", _project("…-0106", 2),
                       source_pdf="a.pdf", placeholder_tags={"P1"}),
        SW.build_entry("MD-522-8110-EG-CV-LAD-0107", _project("…-0107", 5),
                       source_pdf="b.pdf", derived_tags={"GS-01"}),
    ]


class TestShortCodes:
    def test_tail_of_the_drawing_number(self):
        assert SW.short_code("MD-522-8110-EG-CV-LAD-0107") == "0107"

    def test_colliding_serials_are_qualified_by_the_part_that_differs(self):
        """Two areas can share a serial. "LAD-0107" twice would still collide,
        and Excel renames duplicate tabs silently."""
        codes = SW.unique_codes(["MD-522-8110-EG-CV-LAD-0107",
                                 "MD-522-8120-EG-CV-LAD-0107"])
        assert set(codes.values()) == {"8110-0107", "8120-0107"}

    def test_codes_are_always_unique(self):
        codes = SW.unique_codes(["A-1", "A-1", "B-1"])
        assert len(set(codes.values())) == len(set(["A-1", "A-1", "B-1"]))

    def test_sheet_names_fit_excels_limit(self):
        name = SW._sheet_name("8110-0107", "Embedments")
        assert len(name) <= 31

    def test_sheet_names_drop_characters_excel_forbids(self):
        assert "/" not in SW._sheet_name("a/b", "Concrete")


class TestSummaryLayout:
    def _wb(self, tmp_path):
        path = SW.write_set_workbook(_entries(), tmp_path / "set.xlsx",
                                     project_name="Maaden P3")
        return load_workbook(path)

    def test_summary_is_first_and_single(self, tmp_path):
        wb = self._wb(tmp_path)
        assert wb.sheetnames[0] == "Summary"
        assert wb.sheetnames.count("Summary") == 1

    def test_header_is_sl_description_uom_then_one_column_per_drawing(self, tmp_path):
        ws = self._wb(tmp_path)["Summary"]
        header = [ws.cell(row=5, column=c).value for c in range(1, 8)]
        assert header[:3] == ["Sl. #", "Description", "UoM"]
        assert header[3:5] == ["0106", "0107"]
        assert header[5] == "Total"

    def test_full_drawing_numbers_appear_under_the_short_codes(self, tmp_path):
        ws = self._wb(tmp_path)["Summary"]
        assert ws.cell(row=6, column=4).value == "MD-522-8110-EG-CV-LAD-0106"

    def test_quantities_land_under_their_own_drawing(self, tmp_path):
        ws = self._wb(tmp_path)["Summary"]
        rows = [r for r in ws.iter_rows(min_row=7, values_only=True)
                if isinstance(r[0], int) and "Pedestal P1" in str(r[1])]
        assert rows, "pedestal row missing from the summary"
        qty_0106, qty_0107 = rows[0][3], rows[0][4]
        # 5 Nos on 0107 against 2 on 0106, same size — so it must be larger.
        assert qty_0107 > qty_0106

    def test_total_column_sums_across_the_drawings(self, tmp_path):
        ws = self._wb(tmp_path)["Summary"]
        row = next(r for r in ws.iter_rows(min_row=7)
                   if isinstance(r[0].value, int))
        assert str(row[5].value).startswith("=SUM(")

    def test_assumptions_are_named_not_hidden(self, tmp_path):
        ws = self._wb(tmp_path)["Summary"]
        text = " ".join(str(c.value) for r in ws.iter_rows() for c in r
                        if c.value)
        assert "placeholder" in text.lower()

    def test_draft_status_is_stated_on_the_summary(self, tmp_path):
        ws = self._wb(tmp_path)["Summary"]
        assert "UNVERIFIED DRAFT" in str(ws["A3"].value)


class TestDetailSheets:
    def test_named_by_drawing_code_and_activity(self, tmp_path):
        path = SW.write_set_workbook(_entries(), tmp_path / "s.xlsx")
        names = load_workbook(path).sheetnames
        assert "0107_Concrete" in names
        assert any(n.startswith("0106_") for n in names)

    def test_detail_sheet_names_the_full_drawing(self, tmp_path):
        path = SW.write_set_workbook(_entries(), tmp_path / "s.xlsx")
        ws = load_workbook(path)["0107_Concrete"]
        assert "MD-522-8110-EG-CV-LAD-0107" in str(ws["A1"].value)

    def test_an_activity_absent_from_a_drawing_gets_no_sheet(self, tmp_path):
        entries = [SW.build_entry("D-1", _project("D-1", slab=False),
                                  source_pdf="a.pdf")]
        path = SW.write_set_workbook(entries, tmp_path / "s.xlsx")
        assert "D-1_Sump" not in load_workbook(path).sheetnames

    def test_drawing_index_lists_every_sheet(self, tmp_path):
        path = SW.write_set_workbook(_entries(), tmp_path / "s.xlsx")
        ws = load_workbook(path)["Drawings"]
        numbers = [ws.cell(row=r, column=2).value for r in (4, 5)]
        assert numbers == ["MD-522-8110-EG-CV-LAD-0106",
                           "MD-522-8110-EG-CV-LAD-0107"]


class TestEdgeCases:
    def test_a_drawing_with_no_quantities_is_left_out(self, tmp_path):
        empty = Project(project_name="x", drawing_no="D-0")
        entries = _entries() + [SW.build_entry("D-0", empty)]
        path = SW.write_set_workbook(entries, tmp_path / "s.xlsx")
        ws = load_workbook(path)["Summary"]
        header = [ws.cell(row=5, column=c).value for c in range(1, 9)]
        assert "D-0" not in header

    def test_empty_set_still_writes_a_file(self, tmp_path):
        path = SW.write_set_workbook([], tmp_path / "s.xlsx")
        assert path.exists()
        assert load_workbook(path).sheetnames[0] == "Summary"


# ---------------------------------------------------------------------------
# Items the drawings name themselves — see extractors/discovery.py. Five sheets
# in this set carry no pedestal and no slab, so before this block existed they
# reached the workbook with a column and nothing in it.
# ---------------------------------------------------------------------------
DISCOVERY = {
    "items": [
        {"kind": "thickness", "description": "Acid Resistant Epoxy Coating, 4 mm thick",
         "uom": "m2", "qty": None, "basis": "4 mm thick — area not stated",
         "occurrences": 3, "confirm": True, "source": "4MM THK ACID RESISTANT",
         "grid_ref": "J-15 → P-14"},
        {"kind": "rebar", "description": "D25 bars — 16 Nos per element",
         "uom": "Nos", "qty": 16.0, "basis": "16-D25 — bars per element",
         "occurrences": 5, "confirm": False, "source": "16-D25",
         "grid_ref": "I-9 → P-7"},
    ],
    "levels": [],
    "heights": [{"height_m": 2.88, "top": "BOBP EL. 98.380",
                 "bottom": "BOC EL. 95.500"}],
    "grades": [],
    "specifications": [{"value": 35.0, "unit": "MPa",
                        "text": "MINIMUM COMPRESSIVE STRENGTH OF 35 MPa"}],
    "measured": 1,
    "to_confirm": 1,
}


def _discovered_entry(drawing_no="MD-522-8110-EG-CV-LAD-0101"):
    """A details sheet: no pedestal, no slab, plenty specified."""
    project = Project(project_name="Set test", drawing_no=drawing_no)
    return SW.build_entry(drawing_no, project, source_pdf="c.pdf",
                          discovery=DISCOVERY)


class TestDiscoveredItems:
    @pytest.fixture(scope="class")
    def book(self, tmp_path_factory):
        out = tmp_path_factory.mktemp("set") / "SET.xlsx"
        SW.write_set_workbook(_entries() + [_discovered_entry()], out,
                              project_name="Set test")
        return load_workbook(out)

    def test_a_sheet_with_no_pedestal_or_slab_still_earns_a_column(self, book):
        headers = [c.value for c in book["Summary"][5]]
        assert "8110-0101" in headers or "0101" in headers

    def test_discovered_rows_reach_the_summary(self, book):
        text = "\n".join(str(c.value) for row in book["Summary"].iter_rows()
                         for c in row if c.value)
        assert "Acid Resistant Epoxy Coating, 4 mm thick" in text
        assert "D25 bars — 16 Nos per element" in text

    def test_they_are_kept_apart_from_the_priced_rows(self, book):
        text = "\n".join(str(c.value) for row in book["Summary"].iter_rows()
                         for c in row if c.value)
        assert "READ FROM THE DRAWINGS — CONFIRM BEFORE PRICING" in text

    def test_a_per_element_figure_is_not_totalled_across_the_set(self, book):
        """Summing "16 bars per element" across sheets is arithmetic on
        incompatible things, so the Total column says so instead."""
        ws = book["Summary"]
        for row in ws.iter_rows():
            if row[1].value == "D25 bars — 16 Nos per element":
                totals = [c.value for c in row if c.value == "—"]
                assert totals, "the Total cell should be blanked, not summed"
                break
        else:
            pytest.fail("discovered row not found")

    def test_every_discovered_row_is_shaded_as_an_assumption(self, book):
        ws = book["Summary"]
        for row in ws.iter_rows():
            if row[1].value == "Acid Resistant Epoxy Coating, 4 mm thick":
                assert row[3].fill.fgColor.rgb.endswith("FFF2CC")
                break
        else:
            pytest.fail("discovered row not found")

    def test_the_drawing_gets_its_own_items_sheet(self, book):
        assert any(n.endswith("_Drawing items") for n in book.sheetnames)

    def test_that_sheet_keeps_the_text_each_item_came_from(self, book):
        ws = next(book[n] for n in book.sheetnames if n.endswith("_Drawing items"))
        text = "\n".join(str(c.value) for row in ws.iter_rows()
                         for c in row if c.value)
        assert "4MM THK ACID RESISTANT" in text
        assert "J-15 → P-14" in text

    def test_it_records_the_sheet_s_own_specifications_and_implied_heights(self, book):
        ws = next(book[n] for n in book.sheetnames if n.endswith("_Drawing items"))
        text = "\n".join(str(c.value) for row in ws.iter_rows()
                         for c in row if c.value)
        assert "35 MPa" in text
        assert "2.88 m" in text
        assert "BOBP EL. 98.380 over BOC EL. 95.500" in text

    def test_the_index_counts_what_was_read(self, book):
        ws = book["Drawings"]
        headers = [c.value for c in ws[3]]
        assert "Items read" in headers
        col = headers.index("Items read") + 1
        assert any(ws.cell(row=r, column=col).value == 2
                   for r in range(4, ws.max_row + 1))


class TestDoubleCounting:
    """A pedestal callout is read twice: as concrete by the BOM builder, and as
    plan area by the discovery pass. Both belong in the workbook; an estimator
    scanning the two blocks has to be told they are the same element."""

    def test_a_footprint_that_is_already_priced_says_so(self, tmp_path):
        project = _project("…-0106", 2)                  # P1 is 600 x 500
        entry = SW.build_entry(
            "MD-522-8110-EG-CV-LAD-0106", project, source_pdf="a.pdf",
            discovery={"items": [{
                "kind": "box2", "description": "Pedestal 600x500",
                "spec": {"length_mm": 600.0, "width_mm": 500.0},
                "uom": "m2", "qty": 0.3, "basis": "600x500 mm, plan area per unit",
                "occurrences": 2, "confirm": False, "source": "P1(600x500) 2Nos",
                "grid_ref": "D-5 → E-4"}]})
        out = tmp_path / "SET.xlsx"
        SW.write_set_workbook([entry], out)
        ws = load_workbook(out)["Summary"]
        text = "\n".join(str(c.value) for row in ws.iter_rows()
                         for c in row if c.value)
        assert "do not add twice" in text

    def test_an_item_with_no_counterpart_is_not_flagged(self, tmp_path):
        entry = SW.build_entry(
            "MD-522-8110-EG-CV-LAD-0106", _project("…-0106", 2),
            source_pdf="a.pdf",
            discovery={"items": [{
                "kind": "box3", "description": "Manhole 1200x1200x1800",
                "spec": {"length_mm": 1200.0, "width_mm": 1200.0,
                         "height_mm": 1800.0},
                "uom": "m3", "qty": 2.592, "basis": "1200x1200x1800 mm, per unit",
                "occurrences": 1, "confirm": False, "source": "MANHOLE 1200x1200x1800",
                "grid_ref": "A-1"}]})
        out = tmp_path / "SET.xlsx"
        SW.write_set_workbook([entry], out)
        ws = load_workbook(out)["Summary"]
        for row in ws.iter_rows():
            if row[1].value == "Manhole 1200x1200x1800":
                assert "do not add twice" not in str(row[-1].value or "")
                break
        else:
            pytest.fail("discovered row not found")
