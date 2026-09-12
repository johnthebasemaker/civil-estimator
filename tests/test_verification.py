"""Grid references, check print, verification sheet, and the derivation engine.

These are the artefacts a civil team marks up, so the tests care most about the
things that would mislead a checker: a grid reference pointing at the wrong part
of the sheet, or a derived quantity presented as if it were read off the drawing.
"""
from __future__ import annotations

from pathlib import Path

import pytest

from core.derivation import (
    apply_derived, default_rules, derive, rules_from_findings,
)
from core.models import GradeSlab, Pedestal, Project
from extractors.models import (
    ExtractionResult, GradeSlabExtraction, PedestalExtraction, TitleBlockExtraction,
    TranscriptBlock,
)
from extractors.sheet_grid import (
    ALPHABET_WITH_IO, DEFAULT_GRID, SheetGrid, detect_grid,
)
from extractors.verification import (
    VERIFY_HEADERS, build_check_print, collect_items, verification_rows,
)
from extractors.pdf_to_image import open_page

ROOT = Path(__file__).resolve().parent.parent
SAMPLE_PDF = ROOT / "MD-522-8110-EG-CV-LAD-0107_C01.pdf"


# ================= grid =================
class TestSheetGrid:
    def test_column_numbering_runs_right_to_left(self):
        g = DEFAULT_GRID
        assert g.column_number(0.01) == 16
        assert g.column_number(0.99) == 1
        assert g.column_number(0.51) < g.column_number(0.49)

    def test_rows_run_a_to_p_top_to_bottom(self):
        g = DEFAULT_GRID
        assert g.row_letter(0.01) == "A"
        assert g.row_letter(0.99) == "P"

    def test_reference_format(self):
        assert DEFAULT_GRID.ref(0.90, 0.95) == "P-2"

    def test_this_border_includes_i_and_o(self):
        """Most ISO borders skip I and O; this one does not. Assuming the
        convention put every reference below row H off by one."""
        letters = [DEFAULT_GRID.row_letter(i / 16 + 0.001) for i in range(16)]
        assert "I" in letters and "O" in letters
        assert letters == list(ALPHABET_WITH_IO[:16])

    def test_rect_spanning_one_cell_reads_as_one_reference(self):
        assert "→" not in DEFAULT_GRID.ref_for_rect([0.51, 0.51, 0.515, 0.515])

    def test_rect_spanning_cells_shows_a_span(self):
        assert "→" in DEFAULT_GRID.ref_for_rect([0.10, 0.10, 0.50, 0.50])

    def test_empty_rect_gives_no_reference(self):
        assert DEFAULT_GRID.ref_for_rect([]) == ""
        assert DEFAULT_GRID.ref_for_rect(None) == ""

    def test_points_outside_the_sheet_are_clamped(self):
        assert DEFAULT_GRID.ref(-5.0, -5.0) == "A-16"
        assert DEFAULT_GRID.ref(5.0, 5.0) == "P-1"

    @pytest.mark.skipif(not SAMPLE_PDF.exists(), reason="sample drawing missing")
    def test_grid_is_measured_from_the_sheet(self):
        doc, page = open_page(SAMPLE_PDF)
        try:
            g = detect_grid(page)
        finally:
            doc.close()
        assert g.detected
        assert (g.columns, g.rows) == (16, 16)

    def test_detection_falls_back_when_there_is_nothing_to_measure(self):
        class Bare:
            rect = type("R", (), {"width": 100.0, "height": 100.0})()
            rotation_matrix = __import__("fitz").Matrix(1, 0, 0, 1, 0, 0)

            def get_drawings(self):
                return []
        assert detect_grid(Bare()) is DEFAULT_GRID


# ================= verification =================
def _result() -> ExtractionResult:
    r = ExtractionResult(
        source_pdf=str(SAMPLE_PDF), model="qwen2.5vl:7b",
        title_block=TitleBlockExtraction(drawing_no="MD-1", revision="C01",
                                         date="2025-09-30", confidence="high"),
        pedestals=[PedestalExtraction(
            tag="P1", length_mm=600, width_mm=500, quantity=2,
            raw_text="P1(600x500) 2Nos", regex_validated=True,
            confidence="medium", grid_ref="D-7",
            source_rect=[0.53, 0.21, 0.62, 0.24])],
        grade_slabs=[GradeSlabExtraction(tag="GS-01", length_mm=24430,
                                         width_mm=8600, thickness_mm=300)],
        findings={"curb_walls": [{"thickness_mm": 150.0, "height_mm": 150.0,
                                  "raw_text": "CURB WALL 150THK X 150HIGH",
                                  "grid_ref": "B-14",
                                  "source_rect": [0.1, 0.1, 0.2, 0.12]}]},
    )
    return r


class TestVerificationSheet:
    def test_every_extracted_value_becomes_a_check_row(self):
        items = collect_items(_result())
        kinds = {i.kind for i in items}
        assert {"Title block", "Pedestal", "Grade slab", "Curb wall"} <= kinds

    def test_rows_match_the_header_width(self):
        for row in verification_rows(_result()):
            assert len(row) == len(VERIFY_HEADERS)

    def test_signoff_columns_are_left_blank(self):
        for row in verification_rows(_result()):
            assert row[7] == "" and row[8] == "" and row[9] == ""

    def test_pedestal_row_carries_its_callout_and_grid_ref(self):
        row = next(r for r in verification_rows(_result()) if r[1] == "Pedestal")
        assert row[4] == "P1(600x500) 2Nos"
        assert row[5] == "D-7"

    def test_height_is_declared_as_not_on_the_drawing(self):
        item = next(i for i in collect_items(_result()) if i.kind == "Pedestal")
        assert "not on drawing" in item.note.lower()

    def test_review_only_findings_say_they_are_not_in_the_boq(self):
        item = next(i for i in collect_items(_result()) if i.kind == "Curb wall")
        assert "not added to the BOQ" in item.note


@pytest.mark.skipif(not SAMPLE_PDF.exists(), reason="sample drawing missing")
class TestCheckPrint:
    def test_writes_an_image_with_the_located_items_boxed(self, tmp_path):
        doc, page = open_page(SAMPLE_PDF)
        try:
            out = build_check_print(page, _result(), out_path=tmp_path / "chk.png",
                                    width_px=1200)
        finally:
            doc.close()
        assert out.exists() and out.stat().st_size > 10_000
        from PIL import Image
        assert Image.open(out).size[0] == 1200

    def test_items_without_a_location_are_not_boxed(self, tmp_path):
        r = _result()
        for p in r.pedestals:
            p.source_rect = []
        doc, page = open_page(SAMPLE_PDF)
        try:
            out = build_check_print(page, r, out_path=tmp_path / "c.png",
                                    width_px=900)
        finally:
            doc.close()
        assert out.exists()          # must not crash on unlocated items


# ================= derivation =================
def _slab_project() -> Project:
    p = Project(project_name="t", drawing_no="d")
    p.grade_slabs.append(GradeSlab(tag="GS-01", length_m=24.43, width_m=8.6,
                                   thickness_m=0.3))
    return p


class TestDerivation:
    def test_nothing_is_derived_until_a_rule_is_enabled(self):
        assert derive(_slab_project(), default_rules()) == []

    def test_all_rules_produce_valid_core_elements(self):
        rules = default_rules()
        for r in rules:
            r.enabled = True
        p = _slab_project()
        items = derive(p, rules)
        assert len(items) == len(rules)
        apply_derived(p, items)
        from core.bom_builder import build_bom
        assert len(build_bom(p).lines) > 10

    def test_excavation_includes_working_space_both_sides(self):
        rules = [r for r in default_rules() if r.key == "excavation"]
        rules[0].enabled = True
        item = derive(_slab_project(), rules)[0]
        e = item.element
        assert e.length_m == pytest.approx(24.43 + 2 * 0.300)
        assert e.width_m == pytest.approx(8.6 + 2 * 0.300)

    def test_curb_wall_run_is_the_slab_perimeter(self):
        rules = [r for r in default_rules() if r.key == "curb"]
        rules[0].enabled = True
        item = derive(_slab_project(), rules)[0]
        assert item.element.length_m == pytest.approx(2 * (24.43 + 8.6))

    def test_every_derived_item_explains_its_arithmetic(self):
        rules = default_rules()
        for r in rules:
            r.enabled = True
        for item in derive(_slab_project(), rules):
            assert item.explanation and any(ch.isdigit() for ch in item.explanation)

    def test_no_slab_means_nothing_to_derive(self):
        rules = default_rules()
        for r in rules:
            r.enabled = True
        assert derive(Project(project_name="", drawing_no=""), rules) == []

    def test_apply_is_non_destructive(self):
        p = _slab_project()
        p.excavations.append(
            __import__("core.models", fromlist=["Excavation"]).Excavation(
                tag="EXC-GS-01", length_m=1, width_m=1, depth_m=1))
        rules = [r for r in default_rules() if r.key == "excavation"]
        rules[0].enabled = True
        apply_derived(p, derive(p, rules))
        assert len(p.excavations) == 1 and p.excavations[0].length_m == 1

    def test_findings_seed_rule_defaults_without_enabling_them(self):
        rules = rules_from_findings(default_rules(), {
            "curb_walls": [{"thickness_mm": 200.0, "height_mm": 250.0}],
            "epoxy": [{"thickness_mm": 6.0}],
            "thicknesses": [{"thickness_mm": 100.0, "kind": "blinding"}]})
        by = {r.key: r for r in rules}
        assert by["curb"].params["thickness_m"] == 0.200
        assert by["curb"].params["height_m"] == 0.250
        assert by["epoxy"].params["thickness_mm"] == 6.0
        assert by["blinding"].params["thickness_m"] == 0.100
        assert not any(r.enabled for r in rules), "seeding must not enable a rule"

    def test_a_bad_parameter_does_not_kill_the_run(self):
        rules = default_rules()
        for r in rules:
            r.enabled = True
        by = {r.key: r for r in rules}
        by["joints"].params["panel_size_m"] = 0.0      # would divide badly
        assert derive(_slab_project(), rules)          # other rules still produce


# ================= workbook extras =================
def _workbook_with(project: Project, tmp_path: Path) -> Path:
    from core.bom_builder import build_bom
    from core.excel_writer import write_workbook
    return write_workbook(project, build_bom(project), tmp_path / "wb.xlsx")


TEMPLATE_SHEETS = ["Summary", "Assumptions", "Pedestals", "Slab", "Sump",
                   "Joints", "Coating", "Embedments", "Rebar_BBS", "Costing"]


class TestWorkbookExtras:
    def _project(self) -> Project:
        p = _slab_project()
        p.pedestals.append(Pedestal(tag="P1", length_m=0.6, width_m=0.5,
                                    height_m=0.8, quantity=2))
        return p

    def test_extra_sheets_never_disturb_the_mandated_ten(self, tmp_path):
        from extractors import workbook_extras as WE
        path = _workbook_with(self._project(), tmp_path)
        WE.append_verification_sheet(path, _result())
        WE.append_audit_sheet(path, _result())
        rules = default_rules()
        for r in rules:
            r.enabled = True
        WE.append_derivation_sheet(path, derive(self._project(), rules))

        from openpyxl import load_workbook
        names = load_workbook(path).sheetnames
        assert names[:10] == TEMPLATE_SHEETS
        assert names[10:] == ["Verification", "Extraction_Log", "Derivation"]

    def test_appending_twice_does_not_duplicate_sheets(self, tmp_path):
        from extractors import workbook_extras as WE
        path = _workbook_with(self._project(), tmp_path)
        for _ in range(3):
            WE.append_verification_sheet(path, _result())
        from openpyxl import load_workbook
        names = load_workbook(path).sheetnames
        assert names.count("Verification") == 1

    def test_verification_sheet_has_a_row_per_check_item(self, tmp_path):
        from extractors import workbook_extras as WE
        from openpyxl import load_workbook
        path = _workbook_with(self._project(), tmp_path)
        WE.append_verification_sheet(path, _result())
        ws = load_workbook(path)["Verification"]
        headers = [c.value for c in ws[5]]
        assert headers == VERIFY_HEADERS
        body = [r for r in ws.iter_rows(min_row=6, values_only=True)
                if isinstance(r[0], int)]
        assert len(body) == len(verification_rows(_result()))

    def test_derivation_sheet_is_omitted_when_nothing_was_derived(self, tmp_path):
        from extractors import workbook_extras as WE
        from openpyxl import load_workbook
        path = _workbook_with(self._project(), tmp_path)
        WE.append_derivation_sheet(path, [])
        assert "Derivation" not in load_workbook(path).sheetnames


# ================= batch roll-up =================
class TestBatchRollup:
    def test_rollup_sums_across_drawings_and_lists_each(self, tmp_path):
        import run_pipeline as RP

        def one(tag: str, qty: int) -> Project:
            p = Project(project_name="p", drawing_no=tag)
            p.pedestals.append(Pedestal(tag="P1", length_m=0.6, width_m=0.5,
                                        height_m=0.8, quantity=qty))
            return p

        outcomes = [
            RP.RunOutcome(Path("a.pdf"), True, one("DWG-A", 2)),
            RP.RunOutcome(Path("b.pdf"), True, one("DWG-B", 3)),
            RP.RunOutcome(Path("c.pdf"), False, error="boom"),
        ]
        path = RP.write_rollup(outcomes, tmp_path / "roll.xlsx")
        from openpyxl import load_workbook
        wb = load_workbook(path)
        assert wb.sheetnames == ["Set Summary", "Per Drawing"]

        rows = [r for r in wb["Set Summary"].iter_rows(min_row=5, values_only=True)
                if r[0]]
        concrete = [r for r in rows if r[0] == "Structural Concrete"]
        assert concrete, "roll-up must carry concrete quantities"
        # 2 Nos + 3 Nos of the same pedestal, summed across the set
        assert concrete[0][3] == pytest.approx(0.6 * 0.5 * 0.8 * 5 * 1.03, rel=1e-3)

        per = list(wb["Per Drawing"].iter_rows(min_row=2, values_only=True))
        assert [r[1] for r in per] == ["ok", "ok", "FAILED"]
        assert per[2][6] == "boom"

    def test_a_failed_drawing_does_not_stop_the_set(self, tmp_path):
        import run_pipeline as RP
        outcomes = [RP.RunOutcome(Path("x.pdf"), False, error="unreadable")]
        path = RP.write_rollup(outcomes, tmp_path / "r.xlsx")
        assert path.exists()
