"""What the drawing did not say.

The behaviour being defended is a judgement, not a feature: a bill with a
stated gap is more useful than no bill at all. The old code disabled the
button, which meant five sheets in this set produced nothing an estimator could
hold — and the figure that was missing was still missing afterwards.
"""
from __future__ import annotations

import pytest

from core import gaps as G
from core.models import GradeSlab, Pedestal, Project
from extractors.models import ExtractionResult, TitleBlockExtraction

PLACEHOLDER = 1.0


def _project(*, peds=(), slab=False) -> Project:
    p = Project(project_name="t", drawing_no="D-1")
    for tag, height in peds:
        p.pedestals.append(Pedestal(tag=tag, length_m=0.6, width_m=0.5,
                                    height_m=height, quantity=2))
    if slab:
        p.grade_slabs.append(GradeSlab(tag="GS-01", length_m=24.4, width_m=8.6,
                                       thickness_m=0.3))
    return p


def _result(**kw) -> ExtractionResult:
    base = dict(source_pdf="x.pdf",
                title_block=TitleBlockExtraction(drawing_no="D-1"))
    base.update(kw)
    return ExtractionResult(**base)


class TestPlaceholderHeights:
    def test_a_placeholder_height_is_reported_as_an_assumption(self):
        rep = G.report_for(_project(peds=[("P1", PLACEHOLDER)], slab=True),
                           _result(), placeholder_height_m=PLACEHOLDER)
        assumed = [g for g in rep.gaps if g.severity == G.ASSUMED]
        assert any("P1" in g.subject for g in assumed)

    def test_a_real_height_is_not(self):
        rep = G.report_for(_project(peds=[("P1", 0.9)], slab=True), _result(),
                           placeholder_height_m=PLACEHOLDER)
        assert not [g for g in rep.gaps if g.severity == G.ASSUMED]

    def test_it_says_what_the_assumption_moves(self):
        rep = G.report_for(_project(peds=[("P1", PLACEHOLDER)], slab=True),
                           _result(), placeholder_height_m=PLACEHOLDER)
        gap = next(g for g in rep.gaps if g.severity == G.ASSUMED)
        assert "concrete volume" in gap.affects
        assert gap.fix, "a gap with nowhere to fix it is just a complaint"


class TestSheetsWithNothingToPrice:
    def test_a_sections_sheet_is_recorded_rather_than_refused(self):
        result = _result(discovery={"items": [
            {"description": "Acid Resistant, 4 mm thick", "uom": "m2",
             "qty": None, "basis": "area not stated", "source": "4MM THK"}]})
        rep = G.report_for(_project(), result)
        blocking = [g for g in rep.gaps if g.severity == G.BLOCKING]
        assert blocking
        assert any("Drawing_Items" in g.affects for g in blocking)

    def test_an_item_with_no_extent_asks_for_one_figure(self):
        result = _result(discovery={"items": [
            {"description": "Acid Resistant, 4 mm thick", "uom": "m2",
             "qty": None, "basis": "4 mm thick — area not stated",
             "source": "4MM THK ACID RESISTANT"}]})
        rep = G.report_for(_project(), result)
        confirm = [g for g in rep.gaps if g.severity == G.CONFIRM]
        assert any("Acid Resistant" in g.subject for g in confirm)
        assert any("4MM THK ACID RESISTANT" in g.source for g in confirm)

    def test_an_item_with_a_stated_quantity_is_not_a_gap(self):
        result = _result(discovery={"items": [
            {"description": "D25 bars", "uom": "Nos", "qty": 16.0,
             "basis": "16-D25", "source": "16-D25"}]})
        rep = G.report_for(_project(peds=[("P1", 0.9)], slab=True), result)
        assert not [g for g in rep.gaps if g.severity == G.CONFIRM]


class TestMarks:
    def test_marks_are_asked_about_never_priced(self):
        rep = G.report_for(_project(peds=[("P1", 0.9)], slab=True),
                           _result(position_marks={"P1": 19}))
        gap = next(g for g in rep.gaps if g.subject == "Element marks")
        assert gap.severity == G.CONFIRM
        assert "not priced" in gap.affects


class TestRanking:
    def test_assumed_sorts_above_needs_a_figure(self):
        """Assumed is the dangerous class: a number is there and did not come
        off the drawing, so it looks like every other number."""
        result = _result(discovery={"items": [
            {"description": "X", "uom": "m2", "qty": None, "basis": "b",
             "source": "s"}]})
        rep = G.report_for(_project(peds=[("P1", PLACEHOLDER)], slab=True),
                           result, placeholder_height_m=PLACEHOLDER)
        order = [g.severity for g in rep.sorted()]
        assert order.index(G.ASSUMED) < order.index(G.CONFIRM)

    def test_the_headline_says_the_bill_is_usable_and_incomplete(self):
        rep = G.report_for(_project(peds=[("P1", PLACEHOLDER)], slab=True),
                           _result(), placeholder_height_m=PLACEHOLDER)
        assert "usable and incomplete" in rep.headline()

    def test_a_complete_drawing_gets_a_different_headline(self):
        rep = G.report_for(_project(peds=[("P1", 0.9)], slab=True), None)
        assert "usable and incomplete" not in rep.headline()
        assert rep.needs_attention == 0

    def test_every_row_has_somewhere_to_put_the_answer(self):
        result = _result(position_marks={"P1": 4}, discovery={"items": [
            {"description": "X", "uom": "m2", "qty": None, "basis": "b",
             "source": "s"}]})
        rep = G.report_for(_project(peds=[("P1", PLACEHOLDER)], slab=True),
                           result, placeholder_height_m=PLACEHOLDER)
        for gap in rep.gaps:
            if gap.severity in (G.ASSUMED, G.CONFIRM):
                assert gap.fix, f"{gap.subject} says what is wrong but not where"


class TestWorkbookSheet:
    def test_the_sheet_is_written_and_sits_behind_the_summary(self, tmp_path):
        from openpyxl import Workbook, load_workbook

        from extractors import workbook_extras as WE

        path = tmp_path / "wb.xlsx"
        wb = Workbook()
        wb.active.title = "Summary"
        wb.save(path)

        rep = G.report_for(_project(peds=[("P1", PLACEHOLDER)], slab=True),
                           _result(), placeholder_height_m=PLACEHOLDER)
        WE.append_gaps_sheet(path, rep)
        names = load_workbook(path).sheetnames
        assert names[1] == "Gaps_and_Assumptions"

    def test_nothing_is_written_when_there_is_nothing_to_say(self, tmp_path):
        from openpyxl import Workbook, load_workbook

        from extractors import workbook_extras as WE

        path = tmp_path / "wb.xlsx"
        wb = Workbook()
        wb.active.title = "Summary"
        wb.save(path)
        WE.append_gaps_sheet(path, G.GapReport())
        assert load_workbook(path).sheetnames == ["Summary"]
